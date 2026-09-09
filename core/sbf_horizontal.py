"""
Horizontal-format Excel parser and analysis pipeline.
Handles files where each subject is structured in multi-column blocks:
    [Subject Name] [50Hz] [be] [af] [100Hz] [be] [af]
Computes per-person baselines, bucketed after-stimulus values, and paired comparisons.
"""
from __future__ import annotations
import re
from pathlib import Path
from typing import Optional, Union
import numpy as np
import pandas as pd
from scipy import stats as sp_stats
from openpyxl import load_workbook

from .sbf_core import (
    SBFConfig, Recording, ConditionResult,
    filter_signal, buckets_for_recording
)

FREQ_LABEL_RE = re.compile(r"^\s*(\d+)\s*hz\s*$", re.IGNORECASE)
BE_TOKENS = {"be", "before", "pre", "bef"}
AF_TOKENS = {"af", "after", "post", "aft"}


def _is_freq_label(v) -> Optional[int]:
    """Returns the integer Hz value if cell is a frequency label like '50hz', else None."""
    if v is None:
        return None
    s = str(v).strip().lower().replace(" ", "")
    m = FREQ_LABEL_RE.match(s)
    if m:
        return int(m.group(1))
    if s == "50hz":
        return 50
    if s == "100hz":
        return 100
    return None


def _norm_token(v) -> str:
    return str(v).strip().lower() if v is not None and not (isinstance(v, float) and np.isnan(v)) else ""


def parse_horizontal_excel(path: Union[str, Path], sheet: Union[int, str] = 0):
    """
    Parses the header row of an Excel sheet and identifies person-blocks:
        [{name, blocks: {50: {'be': col_1_based, 'af': col_1_based}, 100: ...}}, ...]
    """
    wb = load_workbook(path, data_only=True)
    if isinstance(sheet, int):
        ws = wb.worksheets[sheet]
    else:
        ws = wb[sheet]

    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    persons = []

    c = 0
    while c < len(header):
        h = header[c]
        if h is None or _is_freq_label(h) is not None or _norm_token(h) in BE_TOKENS | AF_TOKENS:
            c += 1
            continue

        name = str(h).strip()
        freq_a = _is_freq_label(header[c + 1]) if c + 1 < len(header) else None
        if freq_a is None:
            c += 1
            continue

        be_a_col, af_a_col = c + 2 + 1, c + 3 + 1  # 1-based index
        freq_b = _is_freq_label(header[c + 4]) if c + 4 < len(header) else None
        be_b_col = af_b_col = None
        if freq_b is not None:
            be_b_col, af_b_col = c + 5 + 1, c + 6 + 1

        ok_a = (_norm_token(header[c + 2]) in BE_TOKENS and
                _norm_token(header[c + 3]) in AF_TOKENS)
        ok_b = freq_b is None or (_norm_token(header[c + 5]) in BE_TOKENS and
                                  _norm_token(header[c + 6]) in AF_TOKENS)

        if not (ok_a and ok_b):
            c += 1
            continue

        block = {"name": name, "blocks": {freq_a: {"be": be_a_col, "af": af_a_col}}}
        if freq_b is not None:
            block["blocks"][freq_b] = {"be": be_b_col, "af": af_b_col}
        persons.append(block)
        c += 7 if freq_b is not None else 4

    return persons, ws, wb


def _load_column(ws, col_idx: int) -> np.ndarray:
    vals = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col_idx).value
        if isinstance(v, (int, float)) and not np.isnan(v):
            vals.append(float(v))
    return np.array(vals, dtype=float)


def parse_horizontal_sheet(df_raw: pd.DataFrame) -> list[dict]:
    """Alternative parser from a pandas DataFrame (e.g. if loaded via pd.read_excel)."""
    header = df_raw.iloc[0].values
    persons = []
    c = 0
    while c < len(header):
        h = header[c]
        if pd.isna(h) or _is_freq_label(h) is not None or _norm_token(h) in BE_TOKENS | AF_TOKENS:
            c += 1
            continue

        name = str(h).strip()
        freq_a = _is_freq_label(header[c + 1]) if c + 1 < len(header) else None
        if freq_a is None:
            c += 1
            continue

        be_a_col, af_a_col = c + 2, c + 3
        freq_b = _is_freq_label(header[c + 4]) if c + 4 < len(header) else None
        be_b_col = af_b_col = None
        if freq_b is not None:
            be_b_col, af_b_col = c + 5, c + 6

        ok_a = (_norm_token(header[c + 2]) in BE_TOKENS and
                _norm_token(header[c + 3]) in AF_TOKENS)
        ok_b = freq_b is None or (_norm_token(header[c + 5]) in BE_TOKENS and
                                  _norm_token(header[c + 6]) in AF_TOKENS)

        if not (ok_a and ok_b):
            c += 1
            continue

        subj = {
            "name": name,
            "recordings": {
                freq_a: {
                    "be": pd.to_numeric(df_raw.iloc[1:, be_a_col], errors="coerce").dropna().values.astype(float),
                    "af": pd.to_numeric(df_raw.iloc[1:, af_a_col], errors="coerce").dropna().values.astype(float),
                }
            }
        }
        if freq_b is not None:
            subj["recordings"][freq_b] = {
                "be": pd.to_numeric(df_raw.iloc[1:, be_b_col], errors="coerce").dropna().values.astype(float),
                "af": pd.to_numeric(df_raw.iloc[1:, af_b_col], errors="coerce").dropna().values.astype(float),
            }
        persons.append(subj)
        c += 7 if freq_b is not None else 4

    return persons


def analyse_horizontal_file(path: Union[str, Path], cfg: SBFConfig,
                            sheet: Union[int, str] = 0,
                            max_buckets: int = 10) -> dict:
    """
    Full pipeline execution on a horizontal Excel file:
    1. Parses subject blocks.
    2. Filters baseline and after-stimulus signals.
    3. Computes per-person baseline statistics and bucket time-courses.
    4. Aggregates group responses per frequency.
    5. Runs paired statistical tests (Paired t-test and Wilcoxon signed-rank).
    """
    person_blocks, ws, wb = parse_horizontal_excel(path, sheet=sheet)
    per_cond: list[ConditionResult] = []
    recordings_full: list[Recording] = []

    try:
        for p in person_blocks:
            name = p["name"]
            for freq, cols in p["blocks"].items():
                be_raw = _load_column(ws, cols["be"])
                af_raw = _load_column(ws, cols["af"])

                be_filt, be_mask = filter_signal(be_raw, cfg)
                af_filt, af_mask = filter_signal(af_raw, cfg)

                recordings_full.append(Recording(person=name, freq_hz=freq, phase="be",
                                                 signal_raw=be_raw, signal_filt=be_filt,
                                                 outlier_mask=be_mask))
                recordings_full.append(Recording(person=name, freq_hz=freq, phase="af",
                                                 signal_raw=af_raw, signal_filt=af_filt,
                                                 outlier_mask=af_mask))

                b_mean = float(np.mean(be_filt)) if len(be_filt) > 0 else 0.0
                b_std = float(np.std(be_filt, ddof=1)) if len(be_filt) > 1 else 0.0
                b_df = buckets_for_recording(af_filt, b_mean, cfg, max_buckets=max_buckets)

                cond = ConditionResult(
                    person=name, freq_hz=freq,
                    baseline_mean=b_mean, baseline_std=b_std,
                    n_baseline_samples=len(be_filt),
                    af_n_samples=len(af_filt),
                    af_outliers_removed=int(af_mask.sum()),
                    af_outlier_pct=100.0 * float(af_mask.sum()) / max(len(af_mask), 1),
                    bucket_df=b_df,
                )
                per_cond.append(cond)
    finally:
        wb.close()

    group_50 = aggregate_group(per_cond, 50, max_buckets)
    group_100 = aggregate_group(per_cond, 100, max_buckets)
    comparison = compare_frequencies(per_cond, max_buckets)

    return {
        "persons": [p["name"] for p in person_blocks],
        "per_cond": per_cond,
        "recordings": recordings_full,
        "group_50": group_50,
        "group_100": group_100,
        "comparison": comparison,
        "config": cfg,
    }


def aggregate_group(per_cond: list[ConditionResult], freq: int,
                    max_buckets: Optional[int] = None) -> pd.DataFrame:
    """Aggregates all subjects for a given frequency."""
    rows_by_freq = [r for r in per_cond if r.freq_hz == freq and len(r.bucket_df) > 0]
    if not rows_by_freq:
        return pd.DataFrame()

    n_buckets = min(len(r.bucket_df) for r in rows_by_freq)
    if max_buckets:
        n_buckets = min(n_buckets, max_buckets)

    out = []
    for k in range(n_buckets):
        ref = rows_by_freq[0].bucket_df.iloc[k]
        means = np.array([r.bucket_df.iloc[k]["mean_SBF"] for r in rows_by_freq])
        changes = np.array([r.bucket_df.iloc[k]["change_pct"] for r in rows_by_freq])
        flagged = np.array([bool(r.bucket_df.iloc[k].get("is_outlier_bucket", False))
                            for r in rows_by_freq])
        keep = ~flagged

        out.append({
            "bucket_idx": k + 1,
            "bucket": ref["bucket"],
            "t_start_s": ref["t_start_s"],
            "t_end_s": ref["t_end_s"],
            "n_people": len(rows_by_freq),
            "n_flagged_excluded": int(flagged.sum()),
            "group_mean_change_pct": float(np.mean(changes)),
            "group_std_change_pct": float(np.std(changes, ddof=1)) if len(changes) > 1 else 0.0,
            "group_sem_change_pct": float(np.std(changes, ddof=1) / np.sqrt(len(changes))) if len(changes) > 1 else 0.0,
            "group_mean_change_pct_clean": float(np.mean(changes[keep])) if keep.sum() > 0 else np.nan,
            "group_std_change_pct_clean": float(np.std(changes[keep], ddof=1)) if keep.sum() > 1 else 0.0,
            "group_mean_SBF": float(np.mean(means)),
            "group_std_SBF": float(np.std(means, ddof=1)) if len(means) > 1 else 0.0,
        })
    return pd.DataFrame(out)


def compare_frequencies(per_cond: list[ConditionResult],
                        max_buckets: Optional[int] = None) -> pd.DataFrame:
    """Computes paired statistical tests (Paired t and Wilcoxon) between 50 Hz and 100 Hz per bucket."""
    by_person: dict[str, dict[int, pd.DataFrame]] = {}
    for r in per_cond:
        by_person.setdefault(r.person, {})[r.freq_hz] = r.bucket_df

    paired_people = [p for p, d in by_person.items() if 50 in d and 100 in d and len(d[50]) > 0 and len(d[100]) > 0]
    if not paired_people:
        return pd.DataFrame()

    min_buckets = min(
        min(len(by_person[p][50]), len(by_person[p][100])) for p in paired_people)
    if max_buckets:
        min_buckets = min(min_buckets, max_buckets)

    out = []
    for k in range(min_buckets):
        v50 = np.array([by_person[p][50].iloc[k]["change_pct"] for p in paired_people])
        v100 = np.array([by_person[p][100].iloc[k]["change_pct"] for p in paired_people])
        diff = v100 - v50

        if len(diff) >= 3:
            try:
                _, p_normal = sp_stats.shapiro(diff)
            except Exception:
                p_normal = 1.0
            t_stat, t_p = sp_stats.ttest_rel(v100, v50, nan_policy="omit")
            try:
                _, w_p = sp_stats.wilcoxon(v100, v50, zero_method="wilcox",
                                           alternative="two-sided", nan_policy="omit")
            except Exception:
                w_p = np.nan
        else:
            t_stat = t_p = w_p = p_normal = np.nan

        out.append({
            "bucket_idx": k + 1,
            "bucket": by_person[paired_people[0]][50].iloc[k]["bucket"],
            "n_paired": len(paired_people),
            "mean_50hz": float(v50.mean()),
            "sd_50hz": float(v50.std(ddof=1)) if len(v50) > 1 else 0.0,
            "mean_100hz": float(v100.mean()),
            "sd_100hz": float(v100.std(ddof=1)) if len(v100) > 1 else 0.0,
            "mean_diff_100minus50": float(diff.mean()),
            "shapiro_p_on_diff": float(p_normal),
            "paired_t_p": float(t_p),
            "wilcoxon_p": float(w_p),
            "preferred_p": float(t_p) if p_normal > 0.05 else float(w_p),
        })
    return pd.DataFrame(out)
