"""
SBF analysis for the horizontal-format Excel file ("flux data 20 subjects.xlsx").

Reads both '12 subjects' and '+ 8 subjects' sheets (combined = 20 subjects),
parses each person's [name][50hz][be][af][100hz][be][af] column block,
filters the raw time-series, computes baselines from 'be', bucketises 'af',
aggregates across subjects, and runs paired 50 Hz vs 100 Hz comparison.

USAGE:
    python run_excel_analysis.py --input "flux data 20 subjects.xlsx" --fs 40
"""
from __future__ import annotations
import argparse, os
from dataclasses import dataclass
import numpy as np
import pandas as pd
from scipy import stats as sp_stats
from scipy.signal import savgol_filter, butter, filtfilt
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ══════════════════════════════════════════════════════════════════════════
# 1. Configuration
# ══════════════════════════════════════════════════════════════════════════
@dataclass
class SBFConfig:
    sampling_rate_hz: float = 40.0
    bucket_seconds: float = 30.0
    use_hampel: bool = True
    hampel_window: int = 21
    hampel_n_sigmas: float = 3.0
    use_savgol: bool = True
    savgol_window: int = 41
    savgol_polyorder: int = 3
    use_lowpass: bool = False
    lowpass_cutoff_hz: float = 2.0
    lowpass_order: int = 4


# ══════════════════════════════════════════════════════════════════════════
# 2. DSP functions
# ══════════════════════════════════════════════════════════════════════════
def hampel_filter(signal: np.ndarray, window_size: int,
                  n_sigmas: float) -> tuple[np.ndarray, np.ndarray]:
    sig = signal.copy()
    mask = np.zeros(len(sig), dtype=bool)
    k = 1.4826
    for i in range(len(sig)):
        lo = max(0, i - window_size)
        hi = min(len(sig), i + window_size + 1)
        window = sig[lo:hi]
        med = np.median(window)
        mad = k * np.median(np.abs(window - med))
        if mad < 1e-12:
            mad = 1e-12
        if np.abs(sig[i] - med) > n_sigmas * mad:
            mask[i] = True
            sig[i] = med
    return sig, mask


def _savitzky_golay(signal: np.ndarray, window: int, polyorder: int) -> np.ndarray:
    return savgol_filter(signal, window, polyorder)


def _butter_lowpass(signal: np.ndarray, fs: float, cutoff: float,
                    order: int) -> np.ndarray:
    nyq = 0.5 * fs
    b, a = butter(order, cutoff / nyq, btype="low")
    return filtfilt(b, a, signal)


def flag_bucket_outliers(means: np.ndarray) -> np.ndarray:
    med = np.median(means)
    mad = 1.4826 * np.median(np.abs(means - med))
    if mad < 1e-12:
        return np.zeros(len(means), dtype=bool)
    stat_flag = np.abs(means - med) > 3.5 * mad
    rel_flag = np.abs(means - med) / max(abs(med), 1e-12) > 0.15
    return stat_flag & rel_flag


# ══════════════════════════════════════════════════════════════════════════
# 3. Excel parsing  —  [name][50hz][be][af][100hz][be][af] layout
# ══════════════════════════════════════════════════════════════════════════
_FREQ_TOKENS = {"50hz", "100hz", "50 hz", "100 hz"}
_BE_TOKENS   = {"be", "before", "pre"}
_AF_TOKENS   = {"af", "after", "post"}


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None and not (isinstance(v, float) and np.isnan(v)) else ""


def _is_freq(v) -> int | None:
    """Returns 50 or 100 if the cell is a frequency label, else None."""
    s = _norm(v).replace(" ", "")
    if s == "50hz": return 50
    if s == "100hz": return 100
    return None


def _extract_column(df_raw: pd.DataFrame, col_idx: int) -> np.ndarray:
    """Pull a numeric column from raw DataFrame (skip header row 0)."""
    series = df_raw.iloc[1:, col_idx]
    vals = pd.to_numeric(series, errors="coerce").dropna().values
    return vals.astype(float)


def parse_horizontal_sheet(df_raw: pd.DataFrame) -> list[dict]:
    """
    Parse one sheet in the horizontal format.
    Returns list of:
        {name, 50: {be: array, af: array}, 100: {be: array, af: array}}
    """
    header = list(df_raw.iloc[0])
    subjects = []
    c = 0
    while c < len(header) - 6:
        name_val = header[c]
        name_s = _norm(name_val)
        # Skip if this cell is a freq/be/af label
        if name_s in _FREQ_TOKENS | _BE_TOKENS | _AF_TOKENS or name_s == "nan" or name_s == "":
            c += 1; continue

        # Expect: name | freqA | be | af | freqB | be | af
        freq_a = _is_freq(header[c + 1]) if c + 1 < len(header) else None
        if freq_a is None:
            c += 1; continue

        be_a_ok = _norm(header[c + 2]) in _BE_TOKENS
        af_a_ok = _norm(header[c + 3]) in _AF_TOKENS
        if not (be_a_ok and af_a_ok):
            c += 1; continue

        freq_b = _is_freq(header[c + 4]) if c + 4 < len(header) else None
        if freq_b is None:
            c += 1; continue

        be_b_ok = _norm(header[c + 5]) in _BE_TOKENS
        af_b_ok = _norm(header[c + 6]) in _AF_TOKENS
        if not (be_b_ok and af_b_ok):
            c += 1; continue

        name = str(name_val).strip()
        be_a = _extract_column(df_raw, c + 2)
        af_a = _extract_column(df_raw, c + 3)
        be_b = _extract_column(df_raw, c + 5)
        af_b = _extract_column(df_raw, c + 6)

        entry = {
            "name": name,
            freq_a: {"be": be_a, "af": af_a},
            freq_b: {"be": be_b, "af": af_b},
        }
        subjects.append(entry)
        c += 7

    return subjects


def load_all_subjects(path: str, sheets: list[str]) -> list[dict]:
    """Load and combine subjects from multiple sheets."""
    all_subjects = []
    seen_names = set()
    for sheet in sheets:
        try:
            df_raw = pd.read_excel(path, sheet_name=sheet, header=None)
        except Exception as e:
            print(f"  WARNING: could not read sheet '{sheet}': {e}")
            continue
        parsed = parse_horizontal_sheet(df_raw)
        for subj in parsed:
            name = subj["name"]
            if name in seen_names:
                print(f"  WARNING: duplicate subject '{name}' in sheet '{sheet}', skipping.")
                continue
            seen_names.add(name)
            all_subjects.append(subj)
        print(f"  Sheet '{sheet}': {len(parsed)} subjects parsed.")
    return all_subjects


# ══════════════════════════════════════════════════════════════════════════
# 4. Signal processing
# ══════════════════════════════════════════════════════════════════════════
def filter_signal(raw: np.ndarray, cfg: SBFConfig) -> tuple[np.ndarray, np.ndarray]:
    sig = raw.copy()
    mask = np.zeros_like(raw, dtype=bool)
    if cfg.use_hampel:
        sig, mask = hampel_filter(sig, cfg.hampel_window, cfg.hampel_n_sigmas)
    if cfg.use_lowpass:
        sig = _butter_lowpass(sig, cfg.sampling_rate_hz,
                              cfg.lowpass_cutoff_hz, cfg.lowpass_order)
    if cfg.use_savgol and len(sig) > cfg.savgol_window:
        sig = _savitzky_golay(sig, cfg.savgol_window, cfg.savgol_polyorder)
    return sig, mask


# ══════════════════════════════════════════════════════════════════════════
# 5. Bucketing
# ══════════════════════════════════════════════════════════════════════════
def buckets_for_recording(af_filt: np.ndarray, baseline_mean: float,
                          cfg: SBFConfig) -> pd.DataFrame:
    fs = cfg.sampling_rate_hz
    bucket_n = int(round(cfg.bucket_seconds * fs))
    rows = []
    start = 0
    while start + bucket_n <= len(af_filt):
        seg = af_filt[start:start + bucket_n]
        t0 = start / fs
        t1 = t0 + cfg.bucket_seconds
        smean = float(seg.mean())
        sstd = float(seg.std(ddof=1))
        change_pct = 100.0 * (smean - baseline_mean) / baseline_mean
        change_ratio = (smean - baseline_mean) / baseline_mean
        slope, _, rval, pval, _ = sp_stats.linregress(
            np.arange(len(seg)) / fs, (seg - baseline_mean) / baseline_mean)
        rows.append({
            "bucket": f"{t0/60:g}-{t1/60:g}",
            "t_start_s": t0, "t_end_s": t1,
            "n_samples": len(seg),
            "mean_SBF": smean, "std_SBF": sstd,
            "change_pct": change_pct, "change_ratio": change_ratio,
            "slope_per_s": slope, "regression_R2": rval**2, "regression_p": pval,
        })
        start += bucket_n
    df = pd.DataFrame(rows)
    if len(df):
        df["is_outlier_bucket"] = flag_bucket_outliers(df["mean_SBF"].values)
    return df


# ══════════════════════════════════════════════════════════════════════════
# 6. Data containers
# ══════════════════════════════════════════════════════════════════════════
@dataclass
class Recording:
    person: str
    freq_hz: int
    phase: str          # "be" or "af"
    signal_raw: np.ndarray
    signal_filt: np.ndarray
    outlier_mask: np.ndarray


@dataclass
class ConditionResult:
    person: str
    freq_hz: int
    baseline_mean: float
    baseline_std: float
    n_baseline_samples: int
    af_n_samples: int
    af_outliers_removed: int
    af_outlier_pct: float
    bucket_df: pd.DataFrame


# ══════════════════════════════════════════════════════════════════════════
# 7. Analysis pipeline
# ══════════════════════════════════════════════════════════════════════════
def analyse(path: str, sheets: list[str], cfg: SBFConfig,
            max_buckets: int | None = 10) -> dict:
    print(f"Loading sheets: {sheets} ...")
    all_subjects = load_all_subjects(path, sheets)
    print(f"Total subjects loaded: {len(all_subjects)}")

    per_cond: list[ConditionResult] = []
    recordings_full: list[Recording] = []

    for subj in all_subjects:
        name = subj["name"]
        for freq in [50, 100]:
            if freq not in subj:
                print(f"  WARNING: {name} missing {freq}Hz data, skipping.")
                continue
            be_raw = subj[freq]["be"]
            af_raw = subj[freq]["af"]
            if len(be_raw) < 10:
                print(f"  WARNING: {name} {freq}Hz 'be' has too few samples ({len(be_raw)}), skipping.")
                continue
            if len(af_raw) < 10:
                print(f"  WARNING: {name} {freq}Hz 'af' has too few samples ({len(af_raw)}), skipping.")
                continue

            be_filt, be_mask = filter_signal(be_raw, cfg)
            af_filt, af_mask = filter_signal(af_raw, cfg)

            baseline_mean = float(be_filt.mean())
            baseline_std = float(be_filt.std(ddof=1)) if len(be_filt) > 1 else 0.0

            bucket_df = buckets_for_recording(af_filt, baseline_mean, cfg)
            if max_buckets:
                bucket_df = bucket_df.head(max_buckets).reset_index(drop=True)

            per_cond.append(ConditionResult(
                person=name, freq_hz=freq,
                baseline_mean=baseline_mean, baseline_std=baseline_std,
                n_baseline_samples=len(be_filt), af_n_samples=len(af_filt),
                af_outliers_removed=int(af_mask.sum()),
                af_outlier_pct=100.0 * af_mask.sum() / max(len(af_raw), 1),
                bucket_df=bucket_df,
            ))
            for phase, raw, filt, mask in [
                ("be", be_raw, be_filt, be_mask),
                ("af", af_raw, af_filt, af_mask),
            ]:
                recordings_full.append(Recording(
                    person=name, freq_hz=freq, phase=phase,
                    signal_raw=raw, signal_filt=filt, outlier_mask=mask,
                ))

    group_50 = _aggregate_group(per_cond, 50, max_buckets)
    group_100 = _aggregate_group(per_cond, 100, max_buckets)
    comparison = _compare_freqs(per_cond, max_buckets)

    return {
        "per_cond": per_cond,
        "recordings_full": recordings_full,
        "group_50": group_50, "group_100": group_100,
        "comparison": comparison,
        "config": cfg,
    }


def _aggregate_group(per_cond: list[ConditionResult], freq: int,
                     max_buckets: int | None) -> pd.DataFrame:
    rows_by_freq = [r for r in per_cond if r.freq_hz == freq]
    if not rows_by_freq:
        return pd.DataFrame()
    n_buckets = min(len(r.bucket_df) for r in rows_by_freq)
    if max_buckets:
        n_buckets = min(n_buckets, max_buckets)

    out = []
    for k in range(n_buckets):
        ref = rows_by_freq[0].bucket_df.iloc[k]
        means = np.array([r.bucket_df.iloc[k]["mean_SBF"] for r in rows_by_freq])
        changes = np.array([r.bucket_df.iloc[k]["change_pct"] for r in rows_by_freq])
        flagged = np.array([bool(r.bucket_df.iloc[k]["is_outlier_bucket"])
                            for r in rows_by_freq])
        keep = ~flagged
        out.append({
            "bucket": ref["bucket"],
            "t_start_s": ref["t_start_s"], "t_end_s": ref["t_end_s"],
            "n_people": len(rows_by_freq),
            "n_flagged_excluded": int(flagged.sum()),
            "group_mean_change_pct": float(np.mean(changes)),
            "group_std_change_pct": float(np.std(changes, ddof=1)) if len(changes) > 1 else 0.0,
            "group_sem_change_pct": float(np.std(changes, ddof=1) / np.sqrt(len(changes))) if len(changes) > 1 else 0.0,
            "group_mean_change_pct_clean": float(np.mean(changes[keep])) if keep.sum() else np.nan,
            "group_std_change_pct_clean": float(np.std(changes[keep], ddof=1)) if keep.sum() > 1 else 0.0,
            "group_mean_SBF": float(np.mean(means)),
            "group_std_SBF_across_people": float(np.std(means, ddof=1)) if len(means) > 1 else 0.0,
        })
    return pd.DataFrame(out)


def _compare_freqs(per_cond: list[ConditionResult],
                   max_buckets: int | None) -> pd.DataFrame:
    by_person: dict[str, dict[int, pd.DataFrame]] = {}
    for r in per_cond:
        by_person.setdefault(r.person, {})[r.freq_hz] = r.bucket_df

    paired_people = [p for p, d in by_person.items() if 50 in d and 100 in d]
    if not paired_people:
        return pd.DataFrame()

    min_buckets = min(
        min(len(by_person[p][50]), len(by_person[p][100])) for p in paired_people)
    if max_buckets:
        min_buckets = min(min_buckets, max_buckets)

    out = []
    for k in range(min_buckets):
        v50 = np.array([by_person[p][50].iloc[k]["change_pct"] for p in paired_people])
        v100 = np.array([by_person[p][100].iloc[k]["change_pct"] for p in paired_people])
        diff = v100 - v50

        if len(diff) >= 3:
            try:
                _, p_normal = sp_stats.shapiro(diff)
            except Exception:
                p_normal = 1.0
            t_stat, t_p = sp_stats.ttest_rel(v100, v50, nan_policy="omit")
            try:
                w_stat, w_p = sp_stats.wilcoxon(v100, v50, zero_method="wilcox",
                                                alternative="two-sided",
                                                nan_policy="omit")
            except Exception:
                w_stat, w_p = np.nan, np.nan
        else:
            t_stat = t_p = w_stat = w_p = p_normal = np.nan

        out.append({
            "bucket": by_person[paired_people[0]][50].iloc[k]["bucket"],
            "n_paired": len(paired_people),
            "mean_50hz": float(v50.mean()), "sd_50hz": float(v50.std(ddof=1)),
            "mean_100hz": float(v100.mean()), "sd_100hz": float(v100.std(ddof=1)),
            "mean_diff_100minus50": float(diff.mean()),
            "shapiro_p_on_diff": float(p_normal),
            "paired_t_p": float(t_p),
            "wilcoxon_p": float(w_p),
            "preferred_p": float(t_p) if p_normal > 0.05 else float(w_p),
        })
    return pd.DataFrame(out)


# ══════════════════════════════════════════════════════════════════════════
# 8. Visualization
# ══════════════════════════════════════════════════════════════════════════
PALETTE = ["#4477AA", "#EE6677", "#228833", "#CCBB44", "#66CCEE", "#AA3377",
           "#000000", "#882255", "#117733", "#999933", "#DDCC77", "#332288",
           "#4477AA", "#EE6677", "#228833", "#CCBB44", "#66CCEE", "#AA3377",
           "#000000", "#882255"]
JOURNAL_RC = {
    "figure.dpi": 110, "savefig.dpi": 250,
    "font.family": "DejaVu Sans", "font.size": 9,
    "axes.titlesize": 10, "axes.labelsize": 9,
    "xtick.labelsize": 8, "ytick.labelsize": 8, "legend.fontsize": 7,
    "axes.linewidth": 0.8, "lines.linewidth": 1.0,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.linestyle": ":", "grid.alpha": 0.4,
}


def plot_group_curves(group_50: pd.DataFrame, group_100: pd.DataFrame,
                      per_cond, output_path: str):
    plt.rcParams.update(JOURNAL_RC)
    fig, axes = plt.subplots(1, 2, figsize=(15, 5), sharey=True)
    for ax, gdf, freq in [(axes[0], group_50, 50), (axes[1], group_100, 100)]:
        if gdf.empty:
            ax.set_title(f"{freq} Hz — no data"); continue
        x = np.arange(len(gdf))
        records = [r for r in per_cond if r.freq_hz == freq]
        for i, r in enumerate(records):
            bdf = r.bucket_df.head(len(gdf))
            ax.plot(x, bdf["change_pct"], "o-",
                    color=PALETTE[i % len(PALETTE)], lw=0.6, ms=3, alpha=0.5,
                    label=r.person)
            flagged = bdf["is_outlier_bucket"].values
            if flagged.any():
                ax.scatter(x[flagged], bdf["change_pct"][flagged],
                           facecolors="none", edgecolors="#C62828",
                           s=80, lw=1.4, zorder=5)
        ax.fill_between(x,
                        gdf["group_mean_change_pct"] - gdf["group_std_change_pct"],
                        gdf["group_mean_change_pct"] + gdf["group_std_change_pct"],
                        color="#37474F", alpha=0.15, label="Group +/- SD")
        ax.plot(x, gdf["group_mean_change_pct"], color="#37474F", lw=2.2,
                marker="s", ms=6, label="Group mean")
        ax.axhline(0, color="black", lw=0.6)
        ax.set_xticks(x)
        ax.set_xticklabels(gdf["bucket"], rotation=0, fontsize=8)
        ax.set_xlabel("Bucket (min after stimulus)")
        ax.set_title(f"{freq} Hz vibration  (n = {gdf['n_people'].iloc[0]} people)",
                     loc="left")
        if freq == 50:
            ax.set_ylabel("Change % vs baseline")
        ax.legend(loc="upper right", ncol=3, fontsize=6)
    fig.suptitle("Recovery response per person — 50 Hz vs 100 Hz", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(output_path, dpi=250, bbox_inches="tight")
    plt.close(fig)


def plot_group_comparison(group_50: pd.DataFrame, group_100: pd.DataFrame,
                          comparison: pd.DataFrame, output_path: str):
    plt.rcParams.update(JOURNAL_RC)
    fig, ax = plt.subplots(figsize=(10, 5))
    x = np.arange(len(group_50))
    w = 0.36
    ax.bar(x - w/2, group_50["group_mean_change_pct"], w,
           yerr=group_50["group_sem_change_pct"], capsize=4,
           color="#4477AA", edgecolor="black", lw=0.6, label="50 Hz")
    ax.bar(x + w/2, group_100["group_mean_change_pct"], w,
           yerr=group_100["group_sem_change_pct"], capsize=4,
           color="#EE6677", edgecolor="black", lw=0.6, label="100 Hz")
    ax.axhline(0, color="black", lw=0.6)
    ax.set_xticks(x)
    ax.set_xticklabels(group_50["bucket"], rotation=0)
    ax.set_xlabel("Bucket (min after stimulus)")
    ax.set_ylabel("Group mean change % +/- SEM")
    n_people = int(group_50["n_people"].iloc[0]) if not group_50.empty else 0
    ax.set_title(f"50 Hz vs 100 Hz  —  group mean (n = {n_people} people)", loc="left")

    if not comparison.empty:
        y_top_50 = (group_50["group_mean_change_pct"] + group_50["group_sem_change_pct"]).max()
        y_top_100 = (group_100["group_mean_change_pct"] + group_100["group_sem_change_pct"]).max()
        y_top = max(y_top_50, y_top_100) * 1.08
        for i, row in comparison.iterrows():
            p = row["preferred_p"]
            sym = "***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "ns"
            if sym != "ns":
                ax.text(i, y_top, sym, ha="center", va="bottom",
                        fontsize=11, fontweight="bold")

    ax.legend()
    ax.text(0.99, -0.18,
            "Significance: *p<.05, **p<.01, ***p<.001 "
            "(paired t-test or Wilcoxon depending on normality)",
            transform=ax.transAxes, ha="right", va="top", fontsize=7, color="#444")
    fig.tight_layout()
    fig.savefig(output_path, dpi=250, bbox_inches="tight")
    plt.close(fig)


def plot_baseline_distribution(per_cond, output_path: str):
    plt.rcParams.update(JOURNAL_RC)
    fig, axes = plt.subplots(1, 2, figsize=(15, 5))
    for ax, freq in zip(axes, [50, 100]):
        rows = [r for r in per_cond if r.freq_hz == freq]
        names = [r.person for r in rows]
        means = np.array([r.baseline_mean for r in rows])
        stds = np.array([r.baseline_std for r in rows])
        order = np.argsort(means)
        x = np.arange(len(rows))
        ax.bar(x, means[order], yerr=stds[order], capsize=3,
               color=PALETTE[0] if freq == 50 else PALETTE[1],
               edgecolor="black", lw=0.5)
        ax.set_xticks(x)
        ax.set_xticklabels([names[i] for i in order],
                           rotation=45, ha="right", fontsize=7)
        ax.set_ylabel("Baseline mean SBF (+/- SD)")
        ax.set_title(f"{freq} Hz — per-person baseline (filtered 'be')", loc="left")
    fig.suptitle("Baseline SBF distributions (20 subjects)", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(output_path, dpi=250, bbox_inches="tight")
    plt.close(fig)


def plot_signals_grid(recordings, output_path: str, freq: int, phase: str,
                      fs: float = 40.0):
    plt.rcParams.update(JOURNAL_RC)
    sel = [r for r in recordings if r.freq_hz == freq and r.phase == phase]
    n = len(sel)
    if n == 0:
        return
    ncols = 4
    nrows = (n + ncols - 1) // ncols
    fig, axes = plt.subplots(nrows, ncols, figsize=(16, 2.5 * nrows), sharex=False)
    axes = np.array(axes).reshape(-1)
    for i, r in enumerate(sel):
        ax = axes[i]
        t = np.arange(len(r.signal_raw)) / fs
        ax.plot(t, r.signal_raw, color="#B0BEC5", lw=0.3, alpha=0.85)
        ax.plot(t, r.signal_filt, color=PALETTE[i % len(PALETTE)], lw=0.8)
        idx = np.where(r.outlier_mask)[0]
        if len(idx):
            ax.scatter(t[idx], r.signal_raw[idx], color="#C62828", s=2, alpha=0.5)
        n_art = int(r.outlier_mask.sum())
        ax.set_title(f"{r.person} ({n_art} outliers)", loc="left", fontsize=7)
        ax.tick_params(axis="both", labelsize=6)
    for j in range(len(sel), len(axes)):
        axes[j].axis("off")
    phase_label = "before (be)" if phase == "be" else "after (af)"
    fig.suptitle(f"{freq} Hz — {phase_label} — raw (grey) vs filtered (color)",
                 fontsize=11)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    fig.savefig(output_path, dpi=250, bbox_inches="tight")
    plt.close(fig)


# ══════════════════════════════════════════════════════════════════════════
# 9. Excel export
# ══════════════════════════════════════════════════════════════════════════
def export_excel(result: dict, output_path: str):
    cfg = result["config"]
    per_cond = result["per_cond"]
    g50 = result["group_50"]
    g100 = result["group_100"]
    comp = result["comparison"]

    wb = Workbook()
    head_fill = PatternFill("solid", start_color="1565C0")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=14, color="0D47A1")
    label_font = Font(bold=True)
    flag_fill = PatternFill("solid", start_color="FFEBEE")
    sig_fill = PatternFill("solid", start_color="E8F5E9")
    thin = Side("thin", color="BBBBBB")
    box = Border(thin, thin, thin, thin)

    def style_header_row(ws, row, n_cols):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row, c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    n_people = len(set(r.person for r in per_cond))

    # ── Summary ──
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Microcirculation SBF analysis — 50 Hz vs 100 Hz vibration (20 subjects)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:E1")
    summary = [
        ("Parameter", "Value"),
        ("Input file", os.path.basename(output_path).replace("_results.xlsx", ".xlsx")),
        ("Sheets combined", "12 subjects, + 8 subjects"),
        ("Sampling rate (Hz)", cfg.sampling_rate_hz),
        ("Bucket length (s)", cfg.bucket_seconds),
        ("Number of people analysed", n_people),
        ("Number of conditions (person x freq)", len(per_cond)),
        ("Frequencies tested", "50 Hz, 100 Hz"),
        ("", ""),
        ("Filter — Hampel window (samples each side)", cfg.hampel_window),
        ("Filter — Hampel threshold (n x MAD)", cfg.hampel_n_sigmas),
        ("Filter — Savitzky-Golay window", cfg.savgol_window),
        ("Filter — Savitzky-Golay polyorder", cfg.savgol_polyorder),
        ("Baseline source", "Mean of filtered 'be' column per person x freq"),
        ("Buckets per group", len(g50)),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws.cell(i, 1, k).font = label_font if i > 3 else head_font
        ws.cell(i, 2, v)
        if i == 3:
            ws.cell(i, 1).fill = head_fill; ws.cell(i, 2).fill = head_fill
            ws.cell(i, 1).font = head_font; ws.cell(i, 2).font = head_font
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 30

    # ── Baselines ──
    ws2 = wb.create_sheet("Baselines")
    headers = ["Person", "Freq Hz", "Baseline mean (filtered)", "Baseline SD",
               "n_baseline_samples", "af samples", "af outliers removed", "af outliers %"]
    for c, h in enumerate(headers, start=1):
        ws2.cell(1, c, h)
    style_header_row(ws2, 1, len(headers))
    for r_i, r in enumerate(per_cond, start=2):
        ws2.cell(r_i, 1, r.person).font = label_font
        ws2.cell(r_i, 2, r.freq_hz)
        ws2.cell(r_i, 3, round(r.baseline_mean, 4))
        ws2.cell(r_i, 4, round(r.baseline_std, 4))
        ws2.cell(r_i, 5, r.n_baseline_samples)
        ws2.cell(r_i, 6, r.af_n_samples)
        ws2.cell(r_i, 7, r.af_outliers_removed)
        ws2.cell(r_i, 8, round(r.af_outlier_pct, 3))
        for c in range(1, len(headers) + 1):
            ws2.cell(r_i, c).border = box
    for col, w in zip("ABCDEFGH", [16, 9, 22, 16, 18, 12, 20, 14]):
        ws2.column_dimensions[col].width = w

    # ── Change% tables ──
    def write_change_table(sheet_name, freq):
        wsX = wb.create_sheet(sheet_name)
        rows_by_freq = [r for r in per_cond if r.freq_hz == freq]
        if not rows_by_freq: return
        n_buckets = min(len(r.bucket_df) for r in rows_by_freq)
        bucket_labels = rows_by_freq[0].bucket_df["bucket"].tolist()[:n_buckets]

        wsX.cell(1, 1, "Person").font = head_font
        wsX.cell(1, 1).fill = head_fill
        for j, b in enumerate(bucket_labels):
            wsX.cell(1, 2 + j, b)
        style_header_row(wsX, 1, 1 + len(bucket_labels))
        for r_i, r in enumerate(rows_by_freq, start=2):
            wsX.cell(r_i, 1, r.person).font = label_font
            for j in range(n_buckets):
                val = r.bucket_df.iloc[j]["change_pct"]
                wsX.cell(r_i, 2 + j, round(float(val), 3))
                if bool(r.bucket_df.iloc[j]["is_outlier_bucket"]):
                    wsX.cell(r_i, 2 + j).fill = flag_fill
                wsX.cell(r_i, 2 + j).border = box
            wsX.cell(r_i, 1).border = box

        gdf = g50 if freq == 50 else g100
        bottom = len(rows_by_freq) + 3
        for lbl, key in [("GROUP MEAN", "group_mean_change_pct"),
                          ("GROUP SD", "group_std_change_pct"),
                          ("GROUP SEM", "group_sem_change_pct"),
                          ("GROUP MEAN (clean)", "group_mean_change_pct_clean")]:
            wsX.cell(bottom, 1, lbl).font = label_font
            for j in range(min(n_buckets, len(gdf))):
                v = gdf.iloc[j][key]
                if not np.isnan(v):
                    wsX.cell(bottom, 2 + j, round(v, 3))
            bottom += 1
        wsX.column_dimensions["A"].width = 22
        for j in range(2, 2 + len(bucket_labels)):
            wsX.column_dimensions[get_column_letter(j)].width = 12

    write_change_table("Change% — 50 Hz", 50)
    write_change_table("Change% — 100 Hz", 100)

    # ── Group Stats ──
    ws5 = wb.create_sheet("Group Stats")
    hdrs = ["Bucket", "n_people",
            "50 Hz: mean chg%", "50 Hz: SD", "50 Hz: SEM", "50 Hz: n flagged",
            "100 Hz: mean chg%", "100 Hz: SD", "100 Hz: SEM", "100 Hz: n flagged"]
    for c, h in enumerate(hdrs, start=1):
        ws5.cell(1, c, h)
    style_header_row(ws5, 1, len(hdrs))
    for r_i in range(len(g50)):
        ws5.cell(r_i + 2, 1, g50.iloc[r_i]["bucket"]).font = label_font
        ws5.cell(r_i + 2, 2, g50.iloc[r_i]["n_people"])
        ws5.cell(r_i + 2, 3, round(g50.iloc[r_i]["group_mean_change_pct"], 3))
        ws5.cell(r_i + 2, 4, round(g50.iloc[r_i]["group_std_change_pct"], 3))
        ws5.cell(r_i + 2, 5, round(g50.iloc[r_i]["group_sem_change_pct"], 3))
        ws5.cell(r_i + 2, 6, int(g50.iloc[r_i]["n_flagged_excluded"]))
        if r_i < len(g100):
            ws5.cell(r_i + 2, 7, round(g100.iloc[r_i]["group_mean_change_pct"], 3))
            ws5.cell(r_i + 2, 8, round(g100.iloc[r_i]["group_std_change_pct"], 3))
            ws5.cell(r_i + 2, 9, round(g100.iloc[r_i]["group_sem_change_pct"], 3))
            ws5.cell(r_i + 2, 10, int(g100.iloc[r_i]["n_flagged_excluded"]))
        for c in range(1, len(hdrs) + 1):
            ws5.cell(r_i + 2, c).border = box
    for col, w in zip("ABCDEFGHIJ", [10, 10, 16, 12, 12, 14, 16, 12, 12, 14]):
        ws5.column_dimensions[col].width = w

    # ── Per-person classic tables ──
    def write_classic_tables(sheet_name, freq):
        wsT = wb.create_sheet(sheet_name)
        rows_by_freq = [r for r in per_cond if r.freq_hz == freq]
        if not rows_by_freq: return
        n_buckets = min(len(r.bucket_df) for r in rows_by_freq)
        BLOCK = 5
        title_fill = PatternFill("solid", start_color="FFE0B2")
        bucket_fill = PatternFill("solid", start_color="FFF8E1")
        std_fill = PatternFill("solid", start_color="FFCCBC")

        for p_i, r in enumerate(rows_by_freq):
            c0 = 1 + p_i * BLOCK
            title = f"{r.person} — {freq} Hz af"
            wsT.cell(1, c0, title).font = Font(bold=True, size=10, color="0D47A1")
            wsT.cell(1, c0).fill = title_fill
            wsT.cell(1, c0).alignment = Alignment(horizontal="center")
            wsT.merge_cells(start_row=1, start_column=c0,
                            end_row=1, end_column=c0 + 3)
            wsT.cell(2, c0, "Bucket").font = head_font
            wsT.cell(2, c0 + 1, "mean ratio").font = head_font
            wsT.cell(2, c0 + 2, "STDEV").font = head_font
            wsT.cell(2, c0 + 3, "std ratio").font = head_font
            for cc in range(c0, c0 + 4):
                wsT.cell(2, cc).fill = head_fill
                wsT.cell(2, cc).alignment = Alignment(horizontal="center",
                                                      wrap_text=True)
            for j in range(n_buckets):
                b = r.bucket_df.iloc[j]
                mean_ratio = float(b["change_ratio"])
                std_ratio = float(b["std_SBF"]) / r.baseline_mean \
                    if r.baseline_mean else float("nan")
                row_r = 3 + j
                wsT.cell(row_r, c0, b["bucket"])
                wsT.cell(row_r, c0).fill = bucket_fill
                wsT.cell(row_r, c0).font = label_font
                wsT.cell(row_r, c0 + 1, round(mean_ratio, 6))
                wsT.cell(row_r, c0 + 2, "STDEV")
                wsT.cell(row_r, c0 + 2).fill = std_fill
                wsT.cell(row_r, c0 + 2).alignment = Alignment(horizontal="center")
                wsT.cell(row_r, c0 + 3, round(std_ratio, 6))
                for cc in range(c0, c0 + 4):
                    wsT.cell(row_r, cc).border = box
                if bool(b["is_outlier_bucket"]):
                    for cc in range(c0, c0 + 4):
                        wsT.cell(row_r, cc).fill = flag_fill

            foot_r = 3 + n_buckets + 1
            wsT.cell(foot_r, c0, "Baseline mean").font = label_font
            wsT.cell(foot_r, c0 + 1, round(r.baseline_mean, 4))
            wsT.cell(foot_r + 1, c0, "Baseline SD").font = label_font
            wsT.cell(foot_r + 1, c0 + 1, round(r.baseline_std, 4))
            wsT.cell(foot_r + 2, c0, "af outliers").font = label_font
            wsT.cell(foot_r + 2, c0 + 1,
                     f"{r.af_outliers_removed} ({r.af_outlier_pct:.1f}%)")

        for p_i in range(len(rows_by_freq)):
            c0 = 1 + p_i * BLOCK
            wsT.column_dimensions[get_column_letter(c0)].width = 9
            wsT.column_dimensions[get_column_letter(c0 + 1)].width = 12
            wsT.column_dimensions[get_column_letter(c0 + 2)].width = 10
            wsT.column_dimensions[get_column_letter(c0 + 3)].width = 12
            wsT.column_dimensions[get_column_letter(c0 + 4)].width = 2

    write_classic_tables("Tables 50Hz (per person)", 50)
    write_classic_tables("Tables 100Hz (per person)", 100)

    # ── 50 vs 100 Hz comparison ──
    ws6 = wb.create_sheet("50 vs 100 Hz")
    hdrs = ["Bucket", "n_paired", "Mean 50 Hz", "SD 50 Hz",
            "Mean 100 Hz", "SD 100 Hz", "Diff (100-50)",
            "Shapiro p", "Paired t p", "Wilcoxon p",
            "Preferred p", "Significance"]
    for c, h in enumerate(hdrs, start=1):
        ws6.cell(1, c, h)
    style_header_row(ws6, 1, len(hdrs))
    for r_i, row in enumerate(comp.itertuples(index=False), start=2):
        ws6.cell(r_i, 1, row.bucket).font = label_font
        ws6.cell(r_i, 2, row.n_paired)
        ws6.cell(r_i, 3, round(row.mean_50hz, 3))
        ws6.cell(r_i, 4, round(row.sd_50hz, 3))
        ws6.cell(r_i, 5, round(row.mean_100hz, 3))
        ws6.cell(r_i, 6, round(row.sd_100hz, 3))
        ws6.cell(r_i, 7, round(row.mean_diff_100minus50, 3))
        ws6.cell(r_i, 8, f"{row.shapiro_p_on_diff:.4f}")
        ws6.cell(r_i, 9, f"{row.paired_t_p:.4f}")
        ws6.cell(r_i, 10, f"{row.wilcoxon_p:.4f}")
        ws6.cell(r_i, 11, f"{row.preferred_p:.4f}")
        p = row.preferred_p
        sym = "***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "ns"
        ws6.cell(r_i, 12, sym)
        if sym != "ns":
            for c in range(1, len(hdrs) + 1):
                ws6.cell(r_i, c).fill = sig_fill
        for c in range(1, len(hdrs) + 1):
            ws6.cell(r_i, c).border = box
    for col, w in zip("ABCDEFGHIJKL", [10, 10, 12, 11, 13, 12, 14, 12, 13, 13, 13, 12]):
        ws6.column_dimensions[col].width = w

    # ── Methodology ──
    ws_m = wb.create_sheet("Methodology")
    notes = [
        "Microcirculation SBF analysis — 50 Hz vs 100 Hz vibration (20 subjects)",
        "",
        "Input format:",
        "  Horizontal Excel layout: [name][50hz][be][af][100hz][be][af] per person.",
        "  Two sheets combined: '12 subjects' + '+ 8 subjects' = 20 subjects total.",
        "  'be' = resting baseline (before vibration); 'af' = recovery (after vibration).",
        "",
        "Per-recording processing (applied to BOTH 'be' and 'af' columns):",
        f"  1. Hampel filter (+/-{cfg.hampel_window} samples, threshold = "
        f"{cfg.hampel_n_sigmas} x MAD) — removes movement-artifact spikes.",
        f"  2. Savitzky-Golay smoother (window={cfg.savgol_window}, "
        f"order={cfg.savgol_polyorder}) — preserves slow physiological waves.",
        "",
        "Baseline (per person x frequency):",
        "  baseline_mean = mean of the FILTERED 'be' column.",
        "",
        "Bucket statistics (per person x frequency, 30-second windows of 'af'):",
        "  change_% = (mean_bucket - baseline_mean) / baseline_mean x 100",
        "  change_ratio = (mean_bucket - baseline_mean) / baseline_mean",
        "",
        "Bucket-level outlier flagging:",
        "  A bucket is flagged if its mean deviates from the overall median by",
        "  both > 3.5 x MAD AND > 15% relative. Flagged buckets are shown in",
        "  light red and excluded from 'clean' group statistics.",
        "",
        "Group aggregation (per frequency, across people):",
        "  group_mean = mean across people of per-person bucket change%",
        "  group_SD   = SD across people  (between-subject variability)",
        "  group_SEM  = SD / sqrt(n)",
        "",
        "Paired 50 vs 100 Hz comparison (per bucket):",
        "  diff = change%_100hz - change%_50hz per person.",
        "  Shapiro-Wilk normality test on diff.",
        "  If normal (p >= .05) -> paired t-test; else -> Wilcoxon signed-rank.",
        "  Significance: * p < .05, ** p < .01, *** p < .001 (uncorrected).",
        "",
        f"  With n = {n_people} people, consider Bonferroni or FDR correction",
        "  before reporting significance in a paper.",
    ]
    for i, line in enumerate(notes, start=1):
        ws_m.cell(i, 1, line)
        if i == 1:
            ws_m.cell(i, 1).font = title_font
    ws_m.column_dimensions["A"].width = 110

    wb.save(output_path)
    print(f"Excel saved: {output_path}")


# ══════════════════════════════════════════════════════════════════════════
# 10. CLI driver
# ══════════════════════════════════════════════════════════════════════════
def main():
    p = argparse.ArgumentParser(
        description="SBF analysis for horizontal-format Excel (20 subjects)")
    p.add_argument("--input", default="flux data 20 subjects.xlsx",
                   help="Path to the Excel file")
    p.add_argument("--sheets", nargs="+",
                   default=["12 subjects", "+ 8 subjects"],
                   help="Sheet names to combine (default: '12 subjects' '+ 8 subjects')")
    p.add_argument("--fs", type=float, default=40.0,
                   help="Data sampling rate in Hz (default 40)")
    p.add_argument("--bucket-seconds", type=float, default=30.0)
    p.add_argument("--max-buckets", type=int, default=10)
    p.add_argument("--hampel-window", type=int, default=21)
    p.add_argument("--hampel-sigmas", type=float, default=3.0)
    p.add_argument("--savgol-window", type=int, default=41)
    p.add_argument("--no-hampel", action="store_true")
    p.add_argument("--no-savgol", action="store_true")
    p.add_argument("--out-prefix", default=None,
                   help="Output file prefix (default: next to input file)")
    p.add_argument("--plot-signals", action="store_true")
    args = p.parse_args()

    input_path = os.path.abspath(args.input)
    if args.out_prefix is None:
        base = os.path.splitext(input_path)[0]
        args.out_prefix = base + "_analysis"

    cfg = SBFConfig(
        sampling_rate_hz=args.fs, bucket_seconds=args.bucket_seconds,
        use_hampel=not args.no_hampel, use_savgol=not args.no_savgol,
        hampel_window=args.hampel_window, hampel_n_sigmas=args.hampel_sigmas,
        savgol_window=args.savgol_window,
    )

    result = analyse(input_path, args.sheets, cfg, max_buckets=args.max_buckets)

    per_cond = result["per_cond"]
    n_people = len(set(r.person for r in per_cond))
    print(f"\nAnalysed {n_people} people, {len(per_cond)} conditions.")

    print(f"\nPer-person baselines:")
    print(f"  {'Person':<16} {'be 50Hz':>10} {'+-SD':>8} {'be 100Hz':>10} {'+-SD':>8} "
          f"{'af out 50':>10} {'af out 100':>11}")
    by_p = {}
    for r in per_cond:
        by_p.setdefault(r.person, {})[r.freq_hz] = r
    for name, conds in sorted(by_p.items()):
        b50 = conds.get(50)
        b100 = conds.get(100)
        print(f"  {name:<16} "
              f"{b50.baseline_mean if b50 else float('nan'):>10.3f} "
              f"{b50.baseline_std  if b50 else float('nan'):>8.3f} "
              f"{b100.baseline_mean if b100 else float('nan'):>10.3f} "
              f"{b100.baseline_std  if b100 else float('nan'):>8.3f} "
              f"{(b50.af_outlier_pct if b50 else float('nan')):>9.2f}% "
              f"{(b100.af_outlier_pct if b100 else float('nan')):>10.2f}%")

    g50, g100 = result["group_50"], result["group_100"]
    print(f"\n=== Group means (change %), 50 Hz ===")
    if not g50.empty:
        print(g50[["bucket", "n_people", "n_flagged_excluded",
                   "group_mean_change_pct", "group_std_change_pct"]].to_string(
            index=False, formatters={
                "group_mean_change_pct": "{:+.2f}".format,
                "group_std_change_pct": "{:.2f}".format,
            }))
    print(f"\n=== Group means (change %), 100 Hz ===")
    if not g100.empty:
        print(g100[["bucket", "n_people", "n_flagged_excluded",
                    "group_mean_change_pct", "group_std_change_pct"]].to_string(
            index=False, formatters={
                "group_mean_change_pct": "{:+.2f}".format,
                "group_std_change_pct": "{:.2f}".format,
            }))

    comp = result["comparison"]
    if not comp.empty:
        print(f"\n=== Paired 50 vs 100 Hz comparison ===")
        print(comp[["bucket", "n_paired", "mean_50hz", "mean_100hz",
                    "mean_diff_100minus50", "paired_t_p", "wilcoxon_p",
                    "preferred_p"]].to_string(index=False, formatters={
                        "mean_50hz": "{:+.2f}".format,
                        "mean_100hz": "{:+.2f}".format,
                        "mean_diff_100minus50": "{:+.2f}".format,
                        "paired_t_p": "{:.4f}".format,
                        "wilcoxon_p": "{:.4f}".format,
                        "preferred_p": "{:.4f}".format,
                    }))

    os.makedirs(os.path.dirname(args.out_prefix) or ".", exist_ok=True)
    xlsx_out = args.out_prefix + "_results.xlsx"
    export_excel(result, xlsx_out)

    print("Generating plots ...")
    plot_group_curves(g50, g100, per_cond,
                      args.out_prefix + "_01_group_curves.png")
    plot_group_comparison(g50, g100, comp,
                          args.out_prefix + "_02_freq_comparison.png")
    plot_baseline_distribution(per_cond,
                               args.out_prefix + "_03_baselines.png")

    if args.plot_signals:
        for freq in [50, 100]:
            for phase in ["be", "af"]:
                plot_signals_grid(result["recordings_full"],
                                  args.out_prefix + f"_signals_{freq}hz_{phase}.png",
                                  freq=freq, phase=phase, fs=cfg.sampling_rate_hz)

    print(f"\nWrote:")
    print(f"  {xlsx_out}")
    print(f"  {args.out_prefix}_01_group_curves.png")
    print(f"  {args.out_prefix}_02_freq_comparison.png")
    print(f"  {args.out_prefix}_03_baselines.png")
    if args.plot_signals:
        print(f"  {args.out_prefix}_signals_*.png")


if __name__ == "__main__":
    main()
