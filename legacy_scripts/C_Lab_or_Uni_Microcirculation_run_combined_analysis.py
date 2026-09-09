"""
Combined SBF analysis for all 26 subjects (20 from Excel + 6 from .txt files).

Computes raw statistics (mean, mode, std, SEM) per person per Hz:
  - Baseline (before): single stats over full filtered signal
  - After (recovery): stats per 30-second bucket

All values in raw perfusion units (no % change).

USAGE:
    python run_combined_analysis.py --dir "c:\\Microcirculation" --excel "flux data 20 subjects.xlsx"
"""
from __future__ import annotations
import argparse, os, re, glob
from dataclasses import dataclass
from typing import Optional
import numpy as np
import pandas as pd
from scipy import stats as sp_stats
from scipy.signal import savgol_filter, butter, filtfilt
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ══════════════════════════════════════════════════════════════════════════
# 1. Configuration
# ══════════════════════════════════════════════════════════════════════════
@dataclass
class SBFConfig:
    sampling_rate_hz: float = 40.0
    bucket_seconds: float = 30.0
    use_hampel: bool = True
    hampel_window: int = 21
    hampel_n_sigmas: float = 3.0
    use_savgol: bool = True
    savgol_window: int = 41
    savgol_polyorder: int = 3
    use_lowpass: bool = False
    lowpass_cutoff_hz: float = 2.0
    lowpass_order: int = 4


# ══════════════════════════════════════════════════════════════════════════
# 2. DSP functions
# ══════════════════════════════════════════════════════════════════════════
def hampel_filter(signal: np.ndarray, window_size: int,
                  n_sigmas: float) -> tuple[np.ndarray, np.ndarray]:
    """Sliding-window Hampel filter. Returns (filtered_signal, outlier_mask)."""
    sig = signal.copy()
    mask = np.zeros(len(sig), dtype=bool)
    k = 1.4826
    for i in range(len(sig)):
        lo = max(0, i - window_size)
        hi = min(len(sig), i + window_size + 1)
        window = sig[lo:hi]
        med = np.median(window)
        mad = k * np.median(np.abs(window - med))
        if mad < 1e-12:
            mad = 1e-12
        if np.abs(sig[i] - med) > n_sigmas * mad:
            mask[i] = True
            sig[i] = med
    return sig, mask


def filter_signal(raw: np.ndarray, cfg: SBFConfig) -> tuple[np.ndarray, np.ndarray]:
    sig = raw.copy()
    mask = np.zeros_like(raw, dtype=bool)
    if cfg.use_hampel:
        sig, mask = hampel_filter(sig, cfg.hampel_window, cfg.hampel_n_sigmas)
    if cfg.use_lowpass:
        nyq = 0.5 * cfg.sampling_rate_hz
        b, a = butter(cfg.lowpass_order, cfg.lowpass_cutoff_hz / nyq, btype="low")
        sig = filtfilt(b, a, sig)
    if cfg.use_savgol and len(sig) > cfg.savgol_window:
        sig = savgol_filter(sig, cfg.savgol_window, cfg.savgol_polyorder)
    return sig, mask


# ══════════════════════════════════════════════════════════════════════════
# 3. Statistics
# ══════════════════════════════════════════════════════════════════════════
def compute_stats(signal: np.ndarray) -> dict:
    """Compute mean, mode (int), std, SEM for a signal array."""
    n = len(signal)
    mean_val = float(np.mean(signal))
    std_val = float(np.std(signal, ddof=1)) if n > 1 else 0.0
    sem_val = std_val / np.sqrt(n) if n > 1 else 0.0
    # Mode: convert to int for meaningful mode calculation
    int_signal = signal.astype(int)
    mode_result = sp_stats.mode(int_signal, keepdims=True)
    mode_val = int(mode_result.mode[0])
    return {
        "mean": round(mean_val, 4),
        "mode": mode_val,
        "std": round(std_val, 4),
        "sem": round(sem_val, 4),
        "n_samples": n,
    }


# ══════════════════════════════════════════════════════════════════════════
# 4. Excel file loader (horizontal format)
# ══════════════════════════════════════════════════════════════════════════
_FREQ_TOKENS = {"50hz", "100hz", "50 hz", "100 hz"}
_BE_TOKENS   = {"be", "before", "pre"}
_AF_TOKENS   = {"af", "after", "post"}


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None and not (isinstance(v, float) and np.isnan(v)) else ""


def _is_freq(v) -> int | None:
    s = _norm(v).replace(" ", "")
    if s == "50hz": return 50
    if s == "100hz": return 100
    return None


def _extract_column(df_raw: pd.DataFrame, col_idx: int) -> np.ndarray:
    series = df_raw.iloc[1:, col_idx]
    vals = pd.to_numeric(series, errors="coerce").dropna().values
    return vals.astype(float)


def parse_horizontal_sheet(df_raw: pd.DataFrame) -> list[dict]:
    header = list(df_raw.iloc[0])
    subjects = []
    c = 0
    while c < len(header) - 6:
        name_val = header[c]
        name_s = _norm(name_val)
        if name_s in _FREQ_TOKENS | _BE_TOKENS | _AF_TOKENS or name_s == "nan" or name_s == "":
            c += 1; continue

        freq_a = _is_freq(header[c + 1]) if c + 1 < len(header) else None
        if freq_a is None:
            c += 1; continue
        be_a_ok = _norm(header[c + 2]) in _BE_TOKENS
        af_a_ok = _norm(header[c + 3]) in _AF_TOKENS
        if not (be_a_ok and af_a_ok):
            c += 1; continue
        freq_b = _is_freq(header[c + 4]) if c + 4 < len(header) else None
        if freq_b is None:
            c += 1; continue
        be_b_ok = _norm(header[c + 5]) in _BE_TOKENS
        af_b_ok = _norm(header[c + 6]) in _AF_TOKENS
        if not (be_b_ok and af_b_ok):
            c += 1; continue

        name = str(name_val).strip()
        be_a = _extract_column(df_raw, c + 2)
        af_a = _extract_column(df_raw, c + 3)
        be_b = _extract_column(df_raw, c + 5)
        af_b = _extract_column(df_raw, c + 6)

        subjects.append({
            "name": name,
            freq_a: {"be": be_a, "af": af_a},
            freq_b: {"be": be_b, "af": af_b},
        })
        c += 7
    return subjects


def load_excel_subjects(path: str, sheets: list[str]) -> list[dict]:
    """Load subjects from Excel. Returns list of {name, source, 50:{be,af}, 100:{be,af}}."""
    results = []
    seen = set()
    for sheet in sheets:
        try:
            df_raw = pd.read_excel(path, sheet_name=sheet, header=None)
        except Exception as e:
            print(f"  WARNING: could not read sheet '{sheet}': {e}")
            continue
        parsed = parse_horizontal_sheet(df_raw)
        for subj in parsed:
            name = subj["name"]
            if name in seen:
                print(f"  WARNING: duplicate '{name}' in sheet '{sheet}', skipping.")
                continue
            seen.add(name)
            subj["source"] = "Excel"
            results.append(subj)
        print(f"  Sheet '{sheet}': {len(parsed)} subjects parsed.")
    return results


# ══════════════════════════════════════════════════════════════════════════
# 5. TXT file loader
# ══════════════════════════════════════════════════════════════════════════
FILE_RE = re.compile(
    r"^(\d{4})\s+(\w+)\s+(\d+)\s*Hz\b.*([AB])\.txt$", re.IGNORECASE
)


def load_channel1(filepath: str) -> Optional[np.ndarray]:
    vals = []
    try:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            for i, line in enumerate(f):
                if i == 0:
                    continue
                parts = line.strip().split("\t")
                if not parts:
                    continue
                try:
                    vals.append(float(parts[0]))
                except (ValueError, IndexError):
                    continue
    except Exception:
        return None
    if len(vals) < 10:
        return None
    return np.array(vals, dtype=float)


def load_txt_subjects(directory: str) -> list[dict]:
    """Load subjects from .txt files. Returns list of {name, source, 50:{be,af}, 100:{be,af}}."""
    # Discover files: person -> freq -> phase -> filepath
    file_map: dict[str, dict[int, dict[str, str]]] = {}
    for path in glob.glob(os.path.join(directory, "*.txt")):
        basename = os.path.basename(path)
        m = FILE_RE.match(basename)
        if not m:
            continue
        _date, person, freq_s, phase = m.groups()
        freq = int(freq_s)
        phase = phase.upper()
        file_map.setdefault(person, {}).setdefault(freq, {})[phase] = path

    results = []
    for person in sorted(file_map.keys()):
        freq_map = file_map[person]
        entry = {"name": person, "source": "TXT"}
        has_data = False
        for freq in [50, 100]:
            if freq not in freq_map:
                continue
            phases = freq_map[freq]
            if "A" not in phases or "B" not in phases:
                print(f"  WARNING: {person} {freq}Hz missing A or B file, skipping this freq.")
                continue
            a_raw = load_channel1(phases["A"])
            b_raw = load_channel1(phases["B"])
            if a_raw is None:
                print(f"  WARNING: {person} {freq}Hz A file corrupted/empty, skipping.")
                continue
            if b_raw is None:
                print(f"  WARNING: {person} {freq}Hz B file corrupted/empty, skipping.")
                continue
            entry[freq] = {"be": a_raw, "af": b_raw}
            has_data = True
        if has_data:
            results.append(entry)
        else:
            print(f"  WARNING: {person} has no valid data, skipping entirely.")
    return results


# ══════════════════════════════════════════════════════════════════════════
# 6. Analysis pipeline
# ══════════════════════════════════════════════════════════════════════════
def analyse_all(subjects: list[dict], cfg: SBFConfig,
                max_buckets: int = 10) -> dict:
    """
    For each subject x freq:
      - Filter baseline -> compute_stats on full signal
      - Filter after -> split into 30-sec buckets -> compute_stats per bucket
    Returns dict with baseline_rows, bucket_rows, subjects list.
    """
    fs = cfg.sampling_rate_hz
    bucket_n = int(round(cfg.bucket_seconds * fs))

    baseline_rows = []
    bucket_rows = []

    for subj in subjects:
        name = subj["name"]
        source = subj.get("source", "?")
        for freq in [50, 100]:
            if freq not in subj:
                continue
            be_raw = subj[freq]["be"]
            af_raw = subj[freq]["af"]
            if len(be_raw) < 10 or len(af_raw) < 10:
                print(f"  WARNING: {name} {freq}Hz too few samples, skipping.")
                continue

            # Filter
            be_filt, be_mask = filter_signal(be_raw, cfg)
            af_filt, af_mask = filter_signal(af_raw, cfg)

            # Baseline stats (full signal)
            be_stats = compute_stats(be_filt)
            baseline_rows.append({
                "name": name, "source": source, "freq_hz": freq,
                **{f"be_{k}": v for k, v in be_stats.items()},
                "be_outliers_removed": int(be_mask.sum()),
            })

            # After: 30-second buckets
            start = 0
            bucket_idx = 0
            while start + bucket_n <= len(af_filt) and bucket_idx < max_buckets:
                seg = af_filt[start:start + bucket_n]
                t0 = start / fs
                t1 = t0 + cfg.bucket_seconds
                time_label = f"{t0/60:g}-{t1/60:g} min"

                seg_stats = compute_stats(seg)
                bucket_rows.append({
                    "name": name, "source": source, "freq_hz": freq,
                    "bucket": bucket_idx + 1,
                    "time_window": time_label,
                    **{f"af_{k}": v for k, v in seg_stats.items()},
                })
                start += bucket_n
                bucket_idx += 1

    return {
        "baseline_rows": baseline_rows,
        "bucket_rows": bucket_rows,
        "subjects": subjects,
        "config": cfg,
    }


# ══════════════════════════════════════════════════════════════════════════
# 7. Excel export
# ══════════════════════════════════════════════════════════════════════════
def export_excel(result: dict, output_path: str):
    cfg = result["config"]
    baseline_rows = result["baseline_rows"]
    bucket_rows = result["bucket_rows"]
    subjects = result["subjects"]

    wb = Workbook()
    head_fill = PatternFill("solid", start_color="1565C0")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=14, color="0D47A1")
    label_font = Font(bold=True)
    txt_fill = PatternFill("solid", start_color="E3F2FD")   # light blue for TXT source
    excel_fill = PatternFill("solid", start_color="FFF3E0")  # light orange for Excel source
    thin = Side("thin", color="BBBBBB")
    box = Border(thin, thin, thin, thin)

    def style_header_row(ws, row, n_cols):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row, c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    # ── Sheet 1: Summary ──
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Combined SBF Analysis — 26 Subjects (Raw Stats)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:E1")

    n_excel = sum(1 for s in subjects if s.get("source") == "Excel")
    n_txt = sum(1 for s in subjects if s.get("source") == "TXT")

    summary = [
        ("Parameter", "Value"),
        ("Total subjects", len(subjects)),
        ("From Excel", f"{n_excel} (flux data 20 subjects.xlsx)"),
        ("From TXT files", f"{n_txt} (individual .txt files)"),
        ("Sampling rate (Hz)", cfg.sampling_rate_hz),
        ("Bucket length (s)", cfg.bucket_seconds),
        ("Frequencies tested", "50 Hz, 100 Hz"),
        ("Statistics computed", "Mean, Mode (int), Std (ddof=1), SEM"),
        ("", ""),
        ("Filter — Hampel window", cfg.hampel_window),
        ("Filter — Hampel threshold (n x MAD)", cfg.hampel_n_sigmas),
        ("Filter — Savitzky-Golay window", cfg.savgol_window),
        ("Filter — Savitzky-Golay polyorder", cfg.savgol_polyorder),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws.cell(i, 1, k).font = label_font if i > 3 else head_font
        ws.cell(i, 2, v)
        if i == 3:
            ws.cell(i, 1).fill = head_fill; ws.cell(i, 2).fill = head_fill
            ws.cell(i, 1).font = head_font; ws.cell(i, 2).font = head_font

    # Subject list
    row_start = len(summary) + 5
    ws.cell(row_start, 1, "Subject").font = head_font
    ws.cell(row_start, 2, "Source").font = head_font
    ws.cell(row_start, 1).fill = head_fill
    ws.cell(row_start, 2).fill = head_fill
    for idx, s in enumerate(subjects):
        r = row_start + 1 + idx
        ws.cell(r, 1, s["name"]).font = label_font
        ws.cell(r, 2, s.get("source", "?"))
        fill = txt_fill if s.get("source") == "TXT" else excel_fill
        ws.cell(r, 1).fill = fill
        ws.cell(r, 2).fill = fill
        ws.cell(r, 1).border = box
        ws.cell(r, 2).border = box

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 35

    # ── Sheet 2: Baseline Stats ──
    ws2 = wb.create_sheet("Baseline Stats")
    headers = ["Name", "Source", "Freq Hz", "Mean", "Mode", "Std", "SEM",
               "N Samples", "Outliers Removed"]
    for c, h in enumerate(headers, start=1):
        ws2.cell(1, c, h)
    style_header_row(ws2, 1, len(headers))

    for r_i, row in enumerate(baseline_rows, start=2):
        ws2.cell(r_i, 1, row["name"]).font = label_font
        ws2.cell(r_i, 2, row["source"])
        ws2.cell(r_i, 3, row["freq_hz"])
        ws2.cell(r_i, 4, row["be_mean"])
        ws2.cell(r_i, 5, row["be_mode"])
        ws2.cell(r_i, 6, row["be_std"])
        ws2.cell(r_i, 7, row["be_sem"])
        ws2.cell(r_i, 8, row["be_n_samples"])
        ws2.cell(r_i, 9, row["be_outliers_removed"])
        fill = txt_fill if row["source"] == "TXT" else excel_fill
        for c in range(1, len(headers) + 1):
            ws2.cell(r_i, c).border = box
            ws2.cell(r_i, c).fill = fill

    for col, w in zip("ABCDEFGHI", [16, 8, 9, 12, 10, 12, 12, 12, 16]):
        ws2.column_dimensions[col].width = w

    # ── Sheet 3: After Stats 50Hz ──
    # ── Sheet 4: After Stats 100Hz ──
    for freq in [50, 100]:
        ws_af = wb.create_sheet(f"After Stats {freq}Hz")
        headers_af = ["Name", "Source", "Bucket", "Time Window",
                      "Mean", "Mode", "Std", "SEM", "N Samples"]
        for c, h in enumerate(headers_af, start=1):
            ws_af.cell(1, c, h)
        style_header_row(ws_af, 1, len(headers_af))

        freq_rows = [r for r in bucket_rows if r["freq_hz"] == freq]
        for r_i, row in enumerate(freq_rows, start=2):
            ws_af.cell(r_i, 1, row["name"]).font = label_font
            ws_af.cell(r_i, 2, row["source"])
            ws_af.cell(r_i, 3, row["bucket"])
            ws_af.cell(r_i, 4, row["time_window"])
            ws_af.cell(r_i, 5, row["af_mean"])
            ws_af.cell(r_i, 6, row["af_mode"])
            ws_af.cell(r_i, 7, row["af_std"])
            ws_af.cell(r_i, 8, row["af_sem"])
            ws_af.cell(r_i, 9, row["af_n_samples"])
            fill = txt_fill if row["source"] == "TXT" else excel_fill
            for c in range(1, len(headers_af) + 1):
                ws_af.cell(r_i, c).border = box
                ws_af.cell(r_i, c).fill = fill

        for col, w in zip("ABCDEFGHI", [16, 8, 8, 14, 12, 10, 12, 12, 12]):
            ws_af.column_dimensions[col].width = w

    # ── Sheet 5: Methodology ──
    ws_m = wb.create_sheet("Methodology")
    notes = [
        "Combined Microcirculation SBF Analysis — 26 Subjects",
        "",
        "Data sources:",
        f"  Excel file: {n_excel} subjects from 'flux data 20 subjects.xlsx'",
        f"  TXT files:  {n_txt} subjects from individual .txt files (Channel 1 only)",
        "",
        "Statistics computed (all in raw perfusion units, NOT % change):",
        "  Mean  — arithmetic mean of filtered signal",
        "  Mode  — most frequent value after converting to integer",
        "  Std   — sample standard deviation (ddof=1)",
        "  SEM   — standard error of the mean (std / sqrt(n))",
        "",
        "Baseline (before vibration):",
        "  Single set of stats computed over the FULL filtered baseline signal.",
        "",
        "After (recovery, post-vibration):",
        f"  Signal split into {cfg.bucket_seconds:.0f}-second buckets ({int(cfg.bucket_seconds * cfg.sampling_rate_hz)} samples each).",
        "  Stats computed per bucket. Max 10 buckets = 5 minutes.",
        "",
        "Signal processing (applied to both baseline and after signals):",
        f"  1. Hampel filter (+/-{cfg.hampel_window} samples, threshold = "
        f"{cfg.hampel_n_sigmas} x MAD) — removes movement-artifact spikes.",
        f"  2. Savitzky-Golay smoother (window={cfg.savgol_window}, "
        f"order={cfg.savgol_polyorder}) — preserves slow physiological waves.",
        "",
        f"Sampling rate: {cfg.sampling_rate_hz} Hz",
    ]
    for i, line in enumerate(notes, start=1):
        ws_m.cell(i, 1, line)
        if i == 1:
            ws_m.cell(i, 1).font = title_font
    ws_m.column_dimensions["A"].width = 110

    wb.save(output_path)
    print(f"Excel saved: {output_path}")


# ══════════════════════════════════════════════════════════════════════════
# 8. Visualization
# ══════════════════════════════════════════════════════════════════════════
PALETTE_EXCEL = "#4477AA"
PALETTE_TXT = "#EE6677"
JOURNAL_RC = {
    "figure.dpi": 110, "savefig.dpi": 250,
    "font.family": "DejaVu Sans", "font.size": 9,
    "axes.titlesize": 10, "axes.labelsize": 9,
    "xtick.labelsize": 8, "ytick.labelsize": 8, "legend.fontsize": 7,
    "axes.linewidth": 0.8, "lines.linewidth": 1.0,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.linestyle": ":", "grid.alpha": 0.4,
}


def plot_baseline_comparison(baseline_rows: list[dict], output_path: str):
    """Bar chart of baseline means per person, grouped by Hz, colored by source."""
    plt.rcParams.update(JOURNAL_RC)
    fig, axes = plt.subplots(1, 2, figsize=(16, 5.5))

    for ax, freq in zip(axes, [50, 100]):
        rows = [r for r in baseline_rows if r["freq_hz"] == freq]
        if not rows:
            ax.set_title(f"{freq} Hz — no data"); continue
        rows = sorted(rows, key=lambda r: r["be_mean"])
        names = [r["name"] for r in rows]
        means = [r["be_mean"] for r in rows]
        stds = [r["be_std"] for r in rows]
        colors = [PALETTE_TXT if r["source"] == "TXT" else PALETTE_EXCEL for r in rows]
        x = np.arange(len(rows))
        ax.bar(x, means, yerr=stds, capsize=3, color=colors,
               edgecolor="black", lw=0.5)
        ax.set_xticks(x)
        ax.set_xticklabels(names, rotation=45, ha="right", fontsize=7)
        ax.set_ylabel("Baseline Mean SBF (perfusion units)")
        ax.set_title(f"{freq} Hz — Baseline per person", loc="left")
        # Legend
        from matplotlib.patches import Patch
        ax.legend(handles=[
            Patch(facecolor=PALETTE_EXCEL, edgecolor="black", label=f"Excel (n={sum(1 for r in rows if r['source']=='Excel')})"),
            Patch(facecolor=PALETTE_TXT, edgecolor="black", label=f"TXT (n={sum(1 for r in rows if r['source']=='TXT')})"),
        ], fontsize=7)

    fig.suptitle("Baseline SBF — 26 Subjects (sorted by mean)", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(output_path, dpi=250, bbox_inches="tight")
    plt.close(fig)


def plot_after_curves(bucket_rows: list[dict], baseline_rows: list[dict],
                      freq: int, output_path: str):
    """Line plot of bucket means over time per person at a given freq."""
    plt.rcParams.update(JOURNAL_RC)
    freq_buckets = [r for r in bucket_rows if r["freq_hz"] == freq]
    if not freq_buckets:
        return

    # Group by person
    by_person: dict[str, list[dict]] = {}
    source_map: dict[str, str] = {}
    for r in freq_buckets:
        by_person.setdefault(r["name"], []).append(r)
        source_map[r["name"]] = r["source"]

    fig, ax = plt.subplots(figsize=(12, 6))
    all_colors = ["#4477AA", "#EE6677", "#228833", "#CCBB44", "#66CCEE", "#AA3377",
                  "#000000", "#882255", "#117733", "#999933", "#DDCC77", "#332288",
                  "#4477AA", "#EE6677", "#228833", "#CCBB44", "#66CCEE", "#AA3377",
                  "#000000", "#882255", "#117733", "#999933", "#DDCC77", "#332288",
                  "#4477AA", "#EE6677"]

    for i, (name, rows) in enumerate(sorted(by_person.items())):
        rows = sorted(rows, key=lambda r: r["bucket"])
        x = [r["bucket"] for r in rows]
        y = [r["af_mean"] for r in rows]
        marker = "s" if source_map[name] == "TXT" else "o"
        ax.plot(x, y, f"{marker}-", color=all_colors[i % len(all_colors)],
                lw=0.8, ms=4, alpha=0.7, label=name)

    ax.set_xlabel("Bucket (30-second window)")
    ax.set_ylabel("Mean SBF (perfusion units)")
    ax.set_title(f"{freq} Hz — After (recovery) per person  |  "
                 f"o = Excel, s = TXT", loc="left")
    ax.legend(loc="upper right", ncol=4, fontsize=6)
    fig.tight_layout()
    fig.savefig(output_path, dpi=250, bbox_inches="tight")
    plt.close(fig)


# ══════════════════════════════════════════════════════════════════════════
# 9. CLI driver
# ══════════════════════════════════════════════════════════════════════════
def main():
    p = argparse.ArgumentParser(
        description="Combined SBF analysis — 26 subjects (Excel + TXT)")
    p.add_argument("--dir", default=".",
                   help="Directory containing .txt data files")
    p.add_argument("--excel", default="flux data 20 subjects.xlsx",
                   help="Path to horizontal Excel file")
    p.add_argument("--sheets", nargs="+",
                   default=["12 subjects", "+ 8 subjects"])
    p.add_argument("--fs", type=float, default=40.0)
    p.add_argument("--bucket-seconds", type=float, default=30.0)
    p.add_argument("--max-buckets", type=int, default=10)
    p.add_argument("--hampel-window", type=int, default=21)
    p.add_argument("--hampel-sigmas", type=float, default=3.0)
    p.add_argument("--savgol-window", type=int, default=41)
    p.add_argument("--no-hampel", action="store_true")
    p.add_argument("--no-savgol", action="store_true")
    p.add_argument("--out-prefix", default=None)
    args = p.parse_args()

    if args.out_prefix is None:
        args.out_prefix = os.path.join(args.dir, "combined_26")

    cfg = SBFConfig(
        sampling_rate_hz=args.fs, bucket_seconds=args.bucket_seconds,
        use_hampel=not args.no_hampel, use_savgol=not args.no_savgol,
        hampel_window=args.hampel_window, hampel_n_sigmas=args.hampel_sigmas,
        savgol_window=args.savgol_window,
    )

    # Load from both sources
    print("=== Loading Excel subjects ===")
    excel_path = os.path.join(args.dir, args.excel) if not os.path.isabs(args.excel) else args.excel
    excel_subjects = load_excel_subjects(excel_path, args.sheets)
    print(f"  Loaded {len(excel_subjects)} subjects from Excel.\n")

    print("=== Loading TXT subjects ===")
    txt_subjects = load_txt_subjects(args.dir)
    print(f"  Loaded {len(txt_subjects)} subjects from TXT files.\n")

    all_subjects = excel_subjects + txt_subjects
    print(f"=== Total: {len(all_subjects)} subjects ===")
    for s in all_subjects:
        freqs = [f for f in [50, 100] if f in s]
        print(f"  {s['name']:<16} [{s.get('source','?'):>5}]  freqs: {freqs}")

    # Run analysis
    print(f"\nRunning analysis (Hampel={cfg.use_hampel}, SavGol={cfg.use_savgol}) ...")
    result = analyse_all(all_subjects, cfg, max_buckets=args.max_buckets)

    baseline_rows = result["baseline_rows"]
    bucket_rows = result["bucket_rows"]

    # Console summary
    print(f"\n=== Baseline Stats (first 10 shown) ===")
    print(f"  {'Name':<16} {'Src':>5} {'Hz':>4} {'Mean':>10} {'Mode':>6} {'Std':>10} {'SEM':>10}")
    for row in baseline_rows[:10]:
        print(f"  {row['name']:<16} {row['source']:>5} {row['freq_hz']:>4} "
              f"{row['be_mean']:>10.2f} {row['be_mode']:>6} "
              f"{row['be_std']:>10.2f} {row['be_sem']:>10.2f}")
    if len(baseline_rows) > 10:
        print(f"  ... and {len(baseline_rows) - 10} more (see Excel)")

    # Write outputs
    os.makedirs(os.path.dirname(args.out_prefix) or ".", exist_ok=True)
    xlsx_out = args.out_prefix + "_analysis.xlsx"
    export_excel(result, xlsx_out)

    print("Generating plots ...")
    plot_baseline_comparison(baseline_rows,
                             args.out_prefix + "_01_baseline_comparison.png")
    plot_after_curves(bucket_rows, baseline_rows, 50,
                      args.out_prefix + "_02_after_curves_50hz.png")
    plot_after_curves(bucket_rows, baseline_rows, 100,
                      args.out_prefix + "_03_after_curves_100hz.png")

    print(f"\nWrote:")
    print(f"  {xlsx_out}")
    print(f"  {args.out_prefix}_01_baseline_comparison.png")
    print(f"  {args.out_prefix}_02_after_curves_50hz.png")
    print(f"  {args.out_prefix}_03_after_curves_100hz.png")
    print("\nDone!")


if __name__ == "__main__":
    main()
