"""
Per-person comparison: Baseline vs After (50Hz & 100Hz) side by side.

For each of 26 subjects, generates:
  - A grouped bar chart: each 30-sec bucket has 3 bars (Baseline, 50Hz After, 100Hz After)
  - Excel sheet with the comparison data

USAGE:
    python run_individual_comparison.py --dir "c:\\Microcirculation" --excel "flux data 20 subjects.xlsx"
"""
from __future__ import annotations
import argparse, os
import numpy as np
import pandas as pd
from scipy import stats as sp_stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Reuse loading and filtering from run_combined_analysis
from run_combined_analysis import (
    SBFConfig, filter_signal, compute_stats,
    load_excel_subjects, load_txt_subjects,
)


# ══════════════════════════════════════════════════════════════════════════
# 1. Per-person analysis
# ══════════════════════════════════════════════════════════════════════════
def analyse_person(subj: dict, cfg: SBFConfig, max_buckets: int = 10) -> dict | None:
    """Analyse one person: baseline + bucketed after for both frequencies.
    Returns dict with name, source, and per-freq data, or None if no data."""
    name = subj["name"]
    source = subj.get("source", "?")
    fs = cfg.sampling_rate_hz
    bucket_n = int(round(cfg.bucket_seconds * fs))

    person_data = {"name": name, "source": source, "freqs": {}}

    for freq in [50, 100]:
        if freq not in subj:
            continue
        be_raw = subj[freq]["be"]
        af_raw = subj[freq]["af"]
        if len(be_raw) < 10 or len(af_raw) < 10:
            continue

        be_filt, _ = filter_signal(be_raw, cfg)
        af_filt, _ = filter_signal(af_raw, cfg)

        be_stats = compute_stats(be_filt)

        buckets = []
        start = 0
        b_idx = 0
        while start + bucket_n <= len(af_filt) and b_idx < max_buckets:
            seg = af_filt[start:start + bucket_n]
            t0 = start / fs
            t1 = t0 + cfg.bucket_seconds
            seg_stats = compute_stats(seg)
            seg_stats["bucket"] = b_idx + 1
            seg_stats["time_window"] = f"{t0/60:g}-{t1/60:g}"
            buckets.append(seg_stats)
            start += bucket_n
            b_idx += 1

        person_data["freqs"][freq] = {
            "baseline": be_stats,
            "buckets": buckets,
        }

    if not person_data["freqs"]:
        return None
    return person_data


# ══════════════════════════════════════════════════════════════════════════
# 2. Per-person bar chart
# ══════════════════════════════════════════════════════════════════════════
JOURNAL_RC = {
    "figure.dpi": 110, "savefig.dpi": 250,
    "font.family": "DejaVu Sans", "font.size": 9,
    "axes.titlesize": 10, "axes.labelsize": 9,
    "xtick.labelsize": 8, "ytick.labelsize": 8, "legend.fontsize": 8,
    "axes.linewidth": 0.8, "lines.linewidth": 1.0,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.linestyle": ":", "grid.alpha": 0.4,
}

COLOR_BASELINE = "#78909C"   # grey-blue
COLOR_50HZ     = "#4477AA"   # blue
COLOR_100HZ    = "#EE6677"   # red


def plot_person(person_data: dict, output_path: str):
    """Bar chart: for each bucket, 3 bars — Baseline, 50Hz After, 100Hz After."""
    plt.rcParams.update(JOURNAL_RC)
    name = person_data["name"]
    source = person_data["source"]
    freqs = person_data["freqs"]

    has_50 = 50 in freqs
    has_100 = 100 in freqs

    # Determine number of buckets (use max across available freqs)
    n_buckets = 0
    if has_50:
        n_buckets = max(n_buckets, len(freqs[50]["buckets"]))
    if has_100:
        n_buckets = max(n_buckets, len(freqs[100]["buckets"]))
    if n_buckets == 0:
        return

    # Get bucket labels from whichever freq has data
    ref_freq = 50 if has_50 else 100
    bucket_labels = [b["time_window"] for b in freqs[ref_freq]["buckets"]]

    # Baselines (average of both if available, otherwise single)
    baseline_50 = freqs[50]["baseline"]["mean"] if has_50 else None
    baseline_100 = freqs[100]["baseline"]["mean"] if has_100 else None
    baseline_50_std = freqs[50]["baseline"]["std"] if has_50 else None
    baseline_100_std = freqs[100]["baseline"]["std"] if has_100 else None

    fig, ax = plt.subplots(figsize=(14, 5.5))

    x = np.arange(n_buckets)
    w = 0.25  # bar width

    # Baseline bars (flat reference, one per bucket position)
    # Show average baseline if both freqs available
    if has_50 and has_100:
        bl_mean = (baseline_50 + baseline_100) / 2
        bl_std = np.sqrt((baseline_50_std**2 + baseline_100_std**2) / 2)
    elif has_50:
        bl_mean = baseline_50
        bl_std = baseline_50_std
    else:
        bl_mean = baseline_100
        bl_std = baseline_100_std

    ax.bar(x - w, [bl_mean] * n_buckets, w, yerr=[bl_std] * n_buckets,
           capsize=3, color=COLOR_BASELINE, edgecolor="black", lw=0.5,
           label=f"Baseline (avg={bl_mean:.1f})")

    # 50Hz after bars
    if has_50:
        means_50 = []
        stds_50 = []
        for k in range(n_buckets):
            if k < len(freqs[50]["buckets"]):
                means_50.append(freqs[50]["buckets"][k]["mean"])
                stds_50.append(freqs[50]["buckets"][k]["std"])
            else:
                means_50.append(0)
                stds_50.append(0)
        ax.bar(x, means_50, w, yerr=stds_50, capsize=3,
               color=COLOR_50HZ, edgecolor="black", lw=0.5,
               label=f"50Hz After")

    # 100Hz after bars
    if has_100:
        means_100 = []
        stds_100 = []
        for k in range(n_buckets):
            if k < len(freqs[100]["buckets"]):
                means_100.append(freqs[100]["buckets"][k]["mean"])
                stds_100.append(freqs[100]["buckets"][k]["std"])
            else:
                means_100.append(0)
                stds_100.append(0)
        ax.bar(x + w, means_100, w, yerr=stds_100, capsize=3,
               color=COLOR_100HZ, edgecolor="black", lw=0.5,
               label=f"100Hz After")

    # Baseline reference lines
    if has_50:
        ax.axhline(baseline_50, color=COLOR_50HZ, lw=1.0, ls="--", alpha=0.5)
        ax.text(n_buckets - 0.5, baseline_50, f"  BL 50Hz={baseline_50:.1f}",
                va="bottom", fontsize=7, color=COLOR_50HZ, alpha=0.8)
    if has_100:
        ax.axhline(baseline_100, color=COLOR_100HZ, lw=1.0, ls="--", alpha=0.5)
        ax.text(n_buckets - 0.5, baseline_100, f"  BL 100Hz={baseline_100:.1f}",
                va="bottom", fontsize=7, color=COLOR_100HZ, alpha=0.8)

    ax.set_xticks(x)
    ax.set_xticklabels(bucket_labels, rotation=0, fontsize=8)
    ax.set_xlabel("Time window (minutes after vibration)")
    ax.set_ylabel("SBF (perfusion units)")
    ax.set_title(f"{name}  [{source}]  —  Baseline vs After (50Hz & 100Hz)",
                 loc="left", fontsize=11, fontweight="bold")
    ax.legend(loc="upper right")

    fig.tight_layout()
    fig.savefig(output_path, dpi=250, bbox_inches="tight")
    plt.close(fig)


# ══════════════════════════════════════════════════════════════════════════
# 3. Excel export — per-person comparison
# ══════════════════════════════════════════════════════════════════════════
def export_comparison_excel(all_person_data: list[dict], cfg: SBFConfig,
                            output_path: str):
    wb = Workbook()
    head_fill = PatternFill("solid", start_color="1565C0")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=13, color="0D47A1")
    label_font = Font(bold=True)
    bl_fill = PatternFill("solid", start_color="ECEFF1")   # grey for baseline
    f50_fill = PatternFill("solid", start_color="E3F2FD")  # blue for 50Hz
    f100_fill = PatternFill("solid", start_color="FFEBEE") # red for 100Hz
    thin = Side("thin", color="BBBBBB")
    box = Border(thin, thin, thin, thin)

    def style_header_row(ws, row, n_cols):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row, c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    # ── Sheet 1: Overview ──
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "Per-Person Comparison: Baseline vs After (50Hz & 100Hz)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:H1")

    headers = ["Name", "Source",
               "BL 50Hz Mean", "BL 50Hz Std",
               "BL 100Hz Mean", "BL 100Hz Std",
               "N Buckets 50Hz", "N Buckets 100Hz"]
    for c, h in enumerate(headers, start=1):
        ws.cell(3, c, h)
    style_header_row(ws, 3, len(headers))

    for r_i, pd_data in enumerate(all_person_data, start=4):
        ws.cell(r_i, 1, pd_data["name"]).font = label_font
        ws.cell(r_i, 2, pd_data["source"])
        if 50 in pd_data["freqs"]:
            ws.cell(r_i, 3, pd_data["freqs"][50]["baseline"]["mean"])
            ws.cell(r_i, 4, pd_data["freqs"][50]["baseline"]["std"])
            ws.cell(r_i, 7, len(pd_data["freqs"][50]["buckets"]))
        if 100 in pd_data["freqs"]:
            ws.cell(r_i, 5, pd_data["freqs"][100]["baseline"]["mean"])
            ws.cell(r_i, 6, pd_data["freqs"][100]["baseline"]["std"])
            ws.cell(r_i, 8, len(pd_data["freqs"][100]["buckets"]))
        for c in range(1, len(headers) + 1):
            ws.cell(r_i, c).border = box
    for col, w in zip("ABCDEFGH", [16, 8, 14, 14, 14, 14, 14, 14]):
        ws.column_dimensions[col].width = w

    # ── Sheet 2: Full comparison table (all persons, all buckets) ──
    ws2 = wb.create_sheet("Full Comparison")
    headers2 = ["Name", "Source", "Bucket", "Time Window",
                "Baseline 50Hz", "After 50Hz Mean", "After 50Hz Mode",
                "After 50Hz Std", "After 50Hz SEM",
                "Baseline 100Hz", "After 100Hz Mean", "After 100Hz Mode",
                "After 100Hz Std", "After 100Hz SEM"]
    for c, h in enumerate(headers2, start=1):
        ws2.cell(1, c, h)
    style_header_row(ws2, 1, len(headers2))

    row_num = 2
    for pd_data in all_person_data:
        name = pd_data["name"]
        source = pd_data["source"]
        freqs = pd_data["freqs"]
        bl_50 = freqs[50]["baseline"]["mean"] if 50 in freqs else None
        bl_100 = freqs[100]["baseline"]["mean"] if 100 in freqs else None

        # Determine max buckets across freqs for this person
        n_b = 0
        if 50 in freqs:
            n_b = max(n_b, len(freqs[50]["buckets"]))
        if 100 in freqs:
            n_b = max(n_b, len(freqs[100]["buckets"]))

        for k in range(n_b):
            ws2.cell(row_num, 1, name).font = label_font
            ws2.cell(row_num, 2, source)
            ws2.cell(row_num, 3, k + 1)

            # Time window label
            tw = None
            if 50 in freqs and k < len(freqs[50]["buckets"]):
                tw = freqs[50]["buckets"][k]["time_window"]
            elif 100 in freqs and k < len(freqs[100]["buckets"]):
                tw = freqs[100]["buckets"][k]["time_window"]
            ws2.cell(row_num, 4, tw or "")

            # 50Hz columns
            if bl_50 is not None:
                ws2.cell(row_num, 5, bl_50)
            if 50 in freqs and k < len(freqs[50]["buckets"]):
                b = freqs[50]["buckets"][k]
                ws2.cell(row_num, 6, b["mean"])
                ws2.cell(row_num, 7, b["mode"])
                ws2.cell(row_num, 8, b["std"])
                ws2.cell(row_num, 9, b["sem"])

            # 100Hz columns
            if bl_100 is not None:
                ws2.cell(row_num, 10, bl_100)
            if 100 in freqs and k < len(freqs[100]["buckets"]):
                b = freqs[100]["buckets"][k]
                ws2.cell(row_num, 11, b["mean"])
                ws2.cell(row_num, 12, b["mode"])
                ws2.cell(row_num, 13, b["std"])
                ws2.cell(row_num, 14, b["sem"])

            for c in range(1, len(headers2) + 1):
                ws2.cell(row_num, c).border = box
            row_num += 1

    for col, w in zip("ABCDEFGHIJKLMN",
                       [16, 8, 8, 12, 14, 14, 12, 12, 12, 14, 14, 12, 12, 12]):
        ws2.column_dimensions[col].width = w

    # ── Per-person sheets ──
    for pd_data in all_person_data:
        name = pd_data["name"]
        source = pd_data["source"]
        freqs = pd_data["freqs"]

        # Sheet name max 31 chars
        sheet_name = name[:28] if len(name) > 28 else name
        ws_p = wb.create_sheet(sheet_name)

        ws_p.cell(1, 1, f"{name} [{source}]").font = title_font
        ws_p.merge_cells("A1:E1")

        # Baseline section
        ws_p.cell(3, 1, "BASELINES").font = Font(bold=True, size=11, color="0D47A1")
        bl_headers = ["Frequency", "Mean", "Mode", "Std", "SEM", "N Samples"]
        for c, h in enumerate(bl_headers, start=1):
            ws_p.cell(4, c, h)
        style_header_row(ws_p, 4, len(bl_headers))

        r = 5
        for freq in [50, 100]:
            if freq not in freqs:
                continue
            bl = freqs[freq]["baseline"]
            fill = f50_fill if freq == 50 else f100_fill
            ws_p.cell(r, 1, f"{freq} Hz").font = label_font
            ws_p.cell(r, 2, bl["mean"])
            ws_p.cell(r, 3, bl["mode"])
            ws_p.cell(r, 4, bl["std"])
            ws_p.cell(r, 5, bl["sem"])
            ws_p.cell(r, 6, bl["n_samples"])
            for c in range(1, len(bl_headers) + 1):
                ws_p.cell(r, c).border = box
                ws_p.cell(r, c).fill = fill
            r += 1

        # After section
        r += 1
        ws_p.cell(r, 1, "AFTER (30-sec buckets)").font = Font(bold=True, size=11, color="0D47A1")
        r += 1

        af_headers = ["Bucket", "Time",
                      "50Hz Mean", "50Hz Mode", "50Hz Std", "50Hz SEM",
                      "100Hz Mean", "100Hz Mode", "100Hz Std", "100Hz SEM"]
        for c, h in enumerate(af_headers, start=1):
            ws_p.cell(r, c, h)
        style_header_row(ws_p, r, len(af_headers))
        r += 1

        n_b = 0
        if 50 in freqs:
            n_b = max(n_b, len(freqs[50]["buckets"]))
        if 100 in freqs:
            n_b = max(n_b, len(freqs[100]["buckets"]))

        for k in range(n_b):
            ws_p.cell(r, 1, k + 1).font = label_font
            tw = None
            if 50 in freqs and k < len(freqs[50]["buckets"]):
                tw = freqs[50]["buckets"][k]["time_window"]
            elif 100 in freqs and k < len(freqs[100]["buckets"]):
                tw = freqs[100]["buckets"][k]["time_window"]
            ws_p.cell(r, 2, tw or "")

            if 50 in freqs and k < len(freqs[50]["buckets"]):
                b = freqs[50]["buckets"][k]
                ws_p.cell(r, 3, b["mean"])
                ws_p.cell(r, 4, b["mode"])
                ws_p.cell(r, 5, b["std"])
                ws_p.cell(r, 6, b["sem"])
                for c in range(3, 7):
                    ws_p.cell(r, c).fill = f50_fill

            if 100 in freqs and k < len(freqs[100]["buckets"]):
                b = freqs[100]["buckets"][k]
                ws_p.cell(r, 7, b["mean"])
                ws_p.cell(r, 8, b["mode"])
                ws_p.cell(r, 9, b["std"])
                ws_p.cell(r, 10, b["sem"])
                for c in range(7, 11):
                    ws_p.cell(r, c).fill = f100_fill

            for c in range(1, len(af_headers) + 1):
                ws_p.cell(r, c).border = box
            r += 1

        for col, w in zip("ABCDEFGHIJ", [8, 12, 12, 10, 12, 12, 12, 10, 12, 12]):
            ws_p.column_dimensions[col].width = w

    wb.save(output_path)
    print(f"Excel saved: {output_path}")


# ══════════════════════════════════════════════════════════════════════════
# 4. CLI
# ══════════════════════════════════════════════════════════════════════════
def main():
    p = argparse.ArgumentParser(
        description="Per-person comparison: Baseline vs After (50Hz & 100Hz)")
    p.add_argument("--dir", default=".",
                   help="Directory containing .txt data files")
    p.add_argument("--excel", default="flux data 20 subjects.xlsx")
    p.add_argument("--sheets", nargs="+",
                   default=["12 subjects", "+ 8 subjects"])
    p.add_argument("--fs", type=float, default=40.0)
    p.add_argument("--bucket-seconds", type=float, default=30.0)
    p.add_argument("--max-buckets", type=int, default=10)
    p.add_argument("--hampel-window", type=int, default=21)
    p.add_argument("--hampel-sigmas", type=float, default=3.0)
    p.add_argument("--savgol-window", type=int, default=41)
    p.add_argument("--no-hampel", action="store_true")
    p.add_argument("--no-savgol", action="store_true")
    p.add_argument("--out-dir", default=None,
                   help="Output directory (default: <dir>/individual_comparison)")
    args = p.parse_args()

    if args.out_dir is None:
        args.out_dir = os.path.join(args.dir, "individual_comparison")
    os.makedirs(args.out_dir, exist_ok=True)

    cfg = SBFConfig(
        sampling_rate_hz=args.fs, bucket_seconds=args.bucket_seconds,
        use_hampel=not args.no_hampel, use_savgol=not args.no_savgol,
        hampel_window=args.hampel_window, hampel_n_sigmas=args.hampel_sigmas,
        savgol_window=args.savgol_window,
    )

    # Load subjects
    print("=== Loading Excel subjects ===")
    excel_path = os.path.join(args.dir, args.excel) if not os.path.isabs(args.excel) else args.excel
    excel_subjects = load_excel_subjects(excel_path, args.sheets)

    print("\n=== Loading TXT subjects ===")
    txt_subjects = load_txt_subjects(args.dir)

    all_subjects = excel_subjects + txt_subjects
    print(f"\n=== Total: {len(all_subjects)} subjects ===\n")

    # Analyse each person
    all_person_data = []
    for subj in all_subjects:
        pd_result = analyse_person(subj, cfg, max_buckets=args.max_buckets)
        if pd_result is None:
            print(f"  {subj['name']}: no valid data, skipping.")
            continue
        all_person_data.append(pd_result)

    print(f"\nAnalysed {len(all_person_data)} subjects.")

    # Generate per-person plots
    print("Generating per-person plots ...")
    for pd_data in all_person_data:
        safe_name = pd_data["name"].replace(" ", "_").replace("/", "_")
        png_path = os.path.join(args.out_dir, f"{safe_name}_comparison.png")
        plot_person(pd_data, png_path)
        print(f"  {png_path}")

    # Export Excel
    xlsx_path = os.path.join(args.out_dir, "individual_comparison.xlsx")
    export_comparison_excel(all_person_data, cfg, xlsx_path)

    print(f"\nDone! {len(all_person_data)} plots + 1 Excel file in: {args.out_dir}")


if __name__ == "__main__":
    main()
