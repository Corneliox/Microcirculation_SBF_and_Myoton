"""
1-Way ANOVA simulation on 26 subjects' after-vibration SBF data.

Splits each 5-min after signal into 3 time segments:
  Seg 1: 0-2 min | Seg 2: 2-3.5 min | Seg 3: 3.5-5 min

Part A: Per-person ANOVA (raw samples within each segment)
Part B: Combined ANOVA (per-person segment means across subjects)
Part C: Simulation — remove 0-6 subjects (313,912 combos), rerun combined ANOVA

USAGE:
    python run_anova_simulation.py --dir "c:\\Microcirculation" --excel "flux data 20 subjects.xlsx"
"""
from __future__ import annotations
import argparse, os, sys
from itertools import combinations
from math import comb
import numpy as np
import pandas as pd
from scipy import stats as sp_stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from run_combined_analysis import (
    SBFConfig, filter_signal, load_excel_subjects, load_txt_subjects,
)

# Segment boundaries in minutes
SEG_BOUNDS = [(0, 2), (2, 3.5), (3.5, 5)]
SEG_LABELS = ["0-2 min", "2-3.5 min", "3.5-5 min"]


# ══════════════════════════════════════════════════════════════════════════
# 1. Segment extraction
# ══════════════════════════════════════════════════════════════════════════
def extract_segments(af_filt: np.ndarray, fs: float = 40.0
                     ) -> list[np.ndarray] | None:
    """Split after signal into 3 time segments. Returns None if too short."""
    total = len(af_filt)
    total_sec = total / fs
    if total_sec < 60:  # need at least 1 min
        return None
    segs = []
    for t0_min, t1_min in SEG_BOUNDS:
        i0 = int(round(t0_min * 60 * fs))
        i1 = int(round(t1_min * 60 * fs))
        i1 = min(i1, total)
        if i0 >= total:
            break
        seg = af_filt[i0:i1]
        if len(seg) < 10:
            break
        segs.append(seg)
    return segs if len(segs) == 3 else None


# ══════════════════════════════════════════════════════════════════════════
# 2. ANOVA helpers
# ══════════════════════════════════════════════════════════════════════════
def compute_eta_squared(groups: list[np.ndarray]) -> float:
    """Compute eta-squared (effect size) for one-way ANOVA."""
    all_data = np.concatenate(groups)
    grand_mean = all_data.mean()
    ss_total = np.sum((all_data - grand_mean) ** 2)
    ss_between = sum(len(g) * (g.mean() - grand_mean) ** 2 for g in groups)
    return float(ss_between / ss_total) if ss_total > 0 else 0.0


def run_anova(groups: list[np.ndarray]) -> dict:
    """Run 1-way ANOVA on list of groups. Returns F, p, eta_sq, sig."""
    if any(len(g) < 2 for g in groups):
        return {"F": np.nan, "p": np.nan, "eta_sq": np.nan, "sig": False}
    F, p = sp_stats.f_oneway(*groups)
    eta_sq = compute_eta_squared(groups)
    return {"F": float(F), "p": float(p), "eta_sq": round(eta_sq, 6),
            "sig": bool(p < 0.05)}


def tukey_posthoc(groups: list[np.ndarray], labels: list[str]) -> list[str]:
    """Run Tukey HSD and return list of significant pair descriptions."""
    try:
        result = sp_stats.tukey_hsd(*groups)
        pairs = []
        for i in range(len(labels)):
            for j in range(i + 1, len(labels)):
                pval = result.pvalue[i][j]
                if pval < 0.05:
                    pairs.append(f"{labels[i]} vs {labels[j]} (p={pval:.4f})")
        return pairs
    except Exception:
        return []


# ══════════════════════════════════════════════════════════════════════════
# 3. Data preparation
# ══════════════════════════════════════════════════════════════════════════
def prepare_data(subjects: list[dict], cfg: SBFConfig
                 ) -> dict[str, dict[int, dict]]:
    """
    For each subject x freq: filter after signal, extract segments,
    compute segment means.

    Returns: {name: {freq: {"segs": [seg1,seg2,seg3], "means": [m1,m2,m3], "source": str}}}
    """
    fs = cfg.sampling_rate_hz
    data = {}
    for subj in subjects:
        name = subj["name"]
        source = subj.get("source", "?")
        data[name] = {}
        for freq in [50, 100]:
            if freq not in subj:
                continue
            af_raw = subj[freq]["af"]
            if len(af_raw) < 10:
                continue
            af_filt, _ = filter_signal(af_raw, cfg)
            segs = extract_segments(af_filt, fs)
            if segs is None:
                print(f"  WARNING: {name} {freq}Hz after signal too short for 3 segments, skipping.")
                continue
            means = [float(s.mean()) for s in segs]
            data[name][freq] = {"segs": segs, "means": means, "source": source}
    return data


# ══════════════════════════════════════════════════════════════════════════
# 4. Part A: Per-person ANOVA
# ══════════════════════════════════════════════════════════════════════════
def per_person_anova(data: dict, freq: int) -> pd.DataFrame:
    """Run ANOVA on raw segments for each person at given freq."""
    rows = []
    for name in sorted(data.keys()):
        if freq not in data[name]:
            continue
        segs = data[name][freq]["segs"]
        source = data[name][freq]["source"]
        res = run_anova(segs)
        tukey = tukey_posthoc(segs, SEG_LABELS) if res["sig"] else []
        rows.append({
            "Name": name, "Source": source, "Freq_Hz": freq,
            "Seg1_mean": round(float(segs[0].mean()), 2),
            "Seg2_mean": round(float(segs[1].mean()), 2),
            "Seg3_mean": round(float(segs[2].mean()), 2),
            "Seg1_n": len(segs[0]), "Seg2_n": len(segs[1]), "Seg3_n": len(segs[2]),
            "F_stat": round(res["F"], 4),
            "p_value": res["p"],
            "eta_squared": res["eta_sq"],
            "Significant": "Yes" if res["sig"] else "No",
            "Tukey_HSD": "; ".join(tukey) if tukey else "",
        })
    return pd.DataFrame(rows)


# ══════════════════════════════════════════════════════════════════════════
# 5. Part B: Combined ANOVA
# ══════════════════════════════════════════════════════════════════════════
def combined_anova(data: dict, freq: int,
                   subject_names: list[str] | None = None) -> dict:
    """
    Combined ANOVA using per-person segment means.
    If subject_names is None, use all subjects with this freq.
    """
    if subject_names is None:
        subject_names = [n for n in sorted(data.keys()) if freq in data[n]]

    seg_means = [[], [], []]
    for name in subject_names:
        if freq not in data[name]:
            continue
        for i, m in enumerate(data[name][freq]["means"]):
            seg_means[i].append(m)

    groups = [np.array(s) for s in seg_means]
    if any(len(g) < 2 for g in groups):
        return {"n_subjects": 0, "F": np.nan, "p": np.nan,
                "eta_sq": np.nan, "sig": False, "tukey": []}

    res = run_anova(groups)
    tukey = tukey_posthoc(groups, SEG_LABELS) if res["sig"] else []
    return {
        "n_subjects": len(groups[0]),
        "group_means": [round(float(g.mean()), 4) for g in groups],
        "group_stds": [round(float(g.std(ddof=1)), 4) for g in groups],
        **res,
        "tukey": tukey,
    }


# ══════════════════════════════════════════════════════════════════════════
# 6. Part C: Simulation loop
# ══════════════════════════════════════════════════════════════════════════
def run_simulation(data: dict, max_remove: int = 6) -> pd.DataFrame:
    """
    For k=0..max_remove: iterate all C(n,k) combos of subjects to remove.
    Run combined ANOVA on remaining subjects for both 50Hz and 100Hz.
    Returns DataFrame with 313,912 rows.
    """
    all_names = sorted(data.keys())
    n = len(all_names)

    # Pre-compute per-person segment means as numpy arrays for speed
    means_50 = {}  # name -> [m1, m2, m3]
    means_100 = {}
    for name in all_names:
        if 50 in data[name]:
            means_50[name] = data[name][50]["means"]
        if 100 in data[name]:
            means_100[name] = data[name][100]["means"]

    names_50 = sorted(means_50.keys())
    names_100 = sorted(means_100.keys())

    # Pre-build arrays for fast indexing
    arr_50 = np.array([means_50[n] for n in names_50])  # (n_50, 3)
    arr_100 = np.array([means_100[n] for n in names_100])  # (n_100, 3)
    idx_50 = {name: i for i, name in enumerate(names_50)}
    idx_100 = {name: i for i, name in enumerate(names_100)}

    total = sum(comb(n, k) for k in range(max_remove + 1))
    print(f"  Simulation: {total:,} combinations to test ...")

    rows = []
    count = 0
    for k in range(max_remove + 1):
        if k == 0:
            combos = [()]
        else:
            combos = combinations(range(n), k)

        for combo in combos:
            removed_names = [all_names[i] for i in combo]
            removed_set = set(removed_names)

            # 50Hz ANOVA
            keep_50 = [i for name, i in idx_50.items() if name not in removed_set]
            if len(keep_50) >= 2:
                sub = arr_50[keep_50]  # (n_keep, 3)
                g1, g2, g3 = sub[:, 0], sub[:, 1], sub[:, 2]
                F_50, p_50 = sp_stats.f_oneway(g1, g2, g3)
                F_50, p_50 = float(F_50), float(p_50)
            else:
                F_50, p_50 = np.nan, np.nan

            # 100Hz ANOVA
            keep_100 = [i for name, i in idx_100.items() if name not in removed_set]
            if len(keep_100) >= 2:
                sub = arr_100[keep_100]
                g1, g2, g3 = sub[:, 0], sub[:, 1], sub[:, 2]
                F_100, p_100 = sp_stats.f_oneway(g1, g2, g3)
                F_100, p_100 = float(F_100), float(p_100)
            else:
                F_100, p_100 = np.nan, np.nan

            rows.append({
                "n_removed": k,
                "removed_subjects": ", ".join(removed_names) if removed_names else "(none)",
                "n_remaining": n - k,
                "F_50hz": round(F_50, 4) if not np.isnan(F_50) else np.nan,
                "p_50hz": p_50,
                "sig_50hz": "Yes" if p_50 < 0.05 else "No",
                "F_100hz": round(F_100, 4) if not np.isnan(F_100) else np.nan,
                "p_100hz": p_100,
                "sig_100hz": "Yes" if p_100 < 0.05 else "No",
            })

            count += 1
            if count % 50000 == 0:
                print(f"    {count:,} / {total:,} done ...")

    print(f"    {count:,} / {total:,} done.")
    return pd.DataFrame(rows)


# ══════════════════════════════════════════════════════════════════════════
# 7. Influence analysis (k=1 subset)
# ══════════════════════════════════════════════════════════════════════════
def influence_analysis(sim_df: pd.DataFrame, baseline_p_50: float,
                       baseline_p_100: float) -> pd.DataFrame:
    """Extract k=1 rows and show each subject's influence."""
    k1 = sim_df[sim_df["n_removed"] == 1].copy()
    k1["removed"] = k1["removed_subjects"]
    k1["delta_p_50hz"] = k1["p_50hz"] - baseline_p_50
    k1["delta_p_100hz"] = k1["p_100hz"] - baseline_p_100
    k1["flips_50hz"] = k1.apply(
        lambda r: "FLIPS" if (baseline_p_50 < 0.05) != (r["p_50hz"] < 0.05) else "", axis=1)
    k1["flips_100hz"] = k1.apply(
        lambda r: "FLIPS" if (baseline_p_100 < 0.05) != (r["p_100hz"] < 0.05) else "", axis=1)
    return k1[["removed", "F_50hz", "p_50hz", "delta_p_50hz", "flips_50hz",
               "F_100hz", "p_100hz", "delta_p_100hz", "flips_100hz"]].reset_index(drop=True)


# ══════════════════════════════════════════════════════════════════════════
# 8. Excel export
# ══════════════════════════════════════════════════════════════════════════
def export_excel(pp_50: pd.DataFrame, pp_100: pd.DataFrame,
                 comb_50: dict, comb_100: dict,
                 sim_df: pd.DataFrame, infl_df: pd.DataFrame,
                 cfg: SBFConfig, output_path: str):
    """Write all results to a multi-sheet Excel file."""
    head_fill = PatternFill("solid", start_color="1565C0")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=14, color="0D47A1")
    label_font = Font(bold=True)
    sig_fill = PatternFill("solid", start_color="E8F5E9")
    ns_fill = PatternFill("solid", start_color="FFEBEE")
    thin = Side("thin", color="BBBBBB")
    box = Border(thin, thin, thin, thin)

    def style_header(ws, row, n_cols):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row, c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    wb = Workbook()

    # ── Summary ──
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "1-Way ANOVA Simulation — 26 Subjects"
    ws["A1"].font = title_font
    ws.merge_cells("A1:D1")
    info = [
        ("Parameter", "Value"),
        ("Subjects", 26),
        ("Frequencies", "50 Hz, 100 Hz (separate)"),
        ("Sampling rate", f"{cfg.sampling_rate_hz} Hz"),
        ("Segment 1", "0-2 min (4800 samples)"),
        ("Segment 2", "2-3.5 min (3600 samples)"),
        ("Segment 3", "3.5-5 min (3600 samples)"),
        ("ANOVA type", "1-way (factor = time segment, 3 levels)"),
        ("Combined ANOVA", "Uses per-person segment means"),
        ("Per-person ANOVA", "Uses raw cleaned samples"),
        ("Simulation", "Remove 0-6 subjects, all C(26,k) combos"),
        ("Total combos", f"{len(sim_df):,}"),
        ("Significance", "p < 0.05"),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(i, 1, k).font = label_font
        ws.cell(i, 2, v)
        if i == 3:
            ws.cell(i, 1).fill = head_fill; ws.cell(i, 2).fill = head_fill
            ws.cell(i, 1).font = head_font; ws.cell(i, 2).font = head_font
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 45

    # ── Per-Person 50Hz ──
    def write_pp_sheet(name, df):
        ws_pp = wb.create_sheet(name)
        headers = list(df.columns)
        for c, h in enumerate(headers, 1):
            ws_pp.cell(1, c, h)
        style_header(ws_pp, 1, len(headers))
        for r_i, row in enumerate(df.itertuples(index=False), start=2):
            for c, val in enumerate(row, start=1):
                cell = ws_pp.cell(r_i, c, val)
                cell.border = box
            sig_val = getattr(row, "Significant", "")
            if sig_val == "Yes":
                for c in range(1, len(headers) + 1):
                    ws_pp.cell(r_i, c).fill = sig_fill
        for c in range(1, len(headers) + 1):
            ws_pp.column_dimensions[get_column_letter(c)].width = 14
        ws_pp.column_dimensions["A"].width = 16
        ws_pp.column_dimensions[get_column_letter(len(headers))].width = 40

    write_pp_sheet("Per-Person 50Hz", pp_50)
    write_pp_sheet("Per-Person 100Hz", pp_100)

    # ── Combined ANOVA ──
    ws_c = wb.create_sheet("Combined ANOVA")
    ws_c["A1"] = "Combined ANOVA (per-person segment means)"
    ws_c["A1"].font = title_font
    ws_c.merge_cells("A1:F1")

    c_headers = ["Freq Hz", "n_subjects",
                 "Seg1 group mean", "Seg2 group mean", "Seg3 group mean",
                 "F_stat", "p_value", "eta_squared", "Significant", "Tukey HSD"]
    for c, h in enumerate(c_headers, 1):
        ws_c.cell(3, c, h)
    style_header(ws_c, 3, len(c_headers))

    for r_i, (freq, res) in enumerate([(50, comb_50), (100, comb_100)], start=4):
        ws_c.cell(r_i, 1, freq).font = label_font
        ws_c.cell(r_i, 2, res["n_subjects"])
        for j, gm in enumerate(res.get("group_means", [np.nan]*3)):
            ws_c.cell(r_i, 3 + j, gm)
        ws_c.cell(r_i, 6, round(res["F"], 4) if not np.isnan(res["F"]) else "N/A")
        ws_c.cell(r_i, 7, f"{res['p']:.6f}" if not np.isnan(res["p"]) else "N/A")
        ws_c.cell(r_i, 8, res["eta_sq"])
        ws_c.cell(r_i, 9, "Yes" if res["sig"] else "No")
        ws_c.cell(r_i, 10, "; ".join(res.get("tukey", [])) or "N/A")
        fill = sig_fill if res["sig"] else ns_fill
        for c in range(1, len(c_headers) + 1):
            ws_c.cell(r_i, c).border = box
            ws_c.cell(r_i, c).fill = fill
    for c, w in enumerate([10, 12, 16, 16, 16, 12, 14, 14, 12, 45], start=1):
        ws_c.column_dimensions[get_column_letter(c)].width = w

    # ── Influence Analysis ──
    ws_i = wb.create_sheet("Influence Analysis")
    ws_i["A1"] = "Influence: effect of removing each individual subject (k=1)"
    ws_i["A1"].font = title_font
    ws_i.merge_cells("A1:I1")
    i_headers = list(infl_df.columns)
    for c, h in enumerate(i_headers, 1):
        ws_i.cell(3, c, h)
    style_header(ws_i, 3, len(i_headers))
    for r_i, row in enumerate(infl_df.itertuples(index=False), start=4):
        for c, val in enumerate(row, start=1):
            if isinstance(val, float):
                val = round(val, 6)
            ws_i.cell(r_i, c, val)
            ws_i.cell(r_i, c).border = box
        # Highlight flips
        if any("FLIPS" in str(getattr(row, col, "")) for col in ["flips_50hz", "flips_100hz"]):
            for c in range(1, len(i_headers) + 1):
                ws_i.cell(r_i, c).fill = PatternFill("solid", start_color="FFF9C4")
    for c, w in enumerate([16, 12, 14, 14, 10, 12, 14, 14, 10], start=1):
        ws_i.column_dimensions[get_column_letter(c)].width = w

    # ── Methodology ──
    ws_m = wb.create_sheet("Methodology")
    notes = [
        "1-Way ANOVA Simulation — Methodology",
        "",
        "Time segments (factor levels for ANOVA):",
        "  Segment 1: 0-2 min after vibration (4800 samples at 40Hz)",
        "  Segment 2: 2-3.5 min after vibration (3600 samples at 40Hz)",
        "  Segment 3: 3.5-5 min after vibration (3600 samples at 40Hz)",
        "",
        "Per-person ANOVA:",
        "  Uses raw cleaned data points within each segment.",
        "  scipy.stats.f_oneway(seg1_raw, seg2_raw, seg3_raw)",
        "  Tests whether SBF levels differ across the 3 time segments for that individual.",
        "",
        "Combined ANOVA:",
        "  Uses per-person segment MEANS (one mean per subject per segment).",
        "  Each subject contributes exactly 1 value per group.",
        "  Tests whether the group of subjects shows a systematic time trend.",
        "",
        "Effect size: eta-squared = SS_between / SS_total",
        "  Small: 0.01, Medium: 0.06, Large: 0.14",
        "",
        "Post-hoc: Tukey HSD (only if ANOVA is significant at p < 0.05)",
        "",
        "Simulation:",
        "  Remove k = 0..6 subjects from the combined ANOVA.",
        "  All C(26,k) combinations are tested exhaustively.",
        "  For each combo: rerun combined ANOVA on remaining subjects.",
        "  Separate results for 50Hz and 100Hz.",
        "",
        "Signal processing (applied before analysis):",
        f"  Hampel filter (window={cfg.hampel_window}, threshold={cfg.hampel_n_sigmas} x MAD)",
        f"  Savitzky-Golay smoother (window={cfg.savgol_window}, order={cfg.savgol_polyorder})",
    ]
    for i, line in enumerate(notes, start=1):
        ws_m.cell(i, 1, line)
        if i == 1:
            ws_m.cell(i, 1).font = title_font
    ws_m.column_dimensions["A"].width = 90

    # ── Simulation Results (large sheet — use ws.append for speed) ──
    print("  Writing simulation results (313k rows) ...")
    ws_sim = wb.create_sheet("Simulation Results")
    sim_headers = list(sim_df.columns)
    for c, h in enumerate(sim_headers, 1):
        ws_sim.cell(1, c, h)
    style_header(ws_sim, 1, len(sim_headers))

    for row_tuple in sim_df.itertuples(index=False):
        ws_sim.append(list(row_tuple))

    for c, w in enumerate([10, 45, 12, 12, 14, 10, 12, 14, 10], start=1):
        ws_sim.column_dimensions[get_column_letter(c)].width = w

    wb.save(output_path)
    print(f"  Excel saved: {output_path}")


# ══════════════════════════════════════════════════════════════════════════
# 9. Plots
# ══════════════════════════════════════════════════════════════════════════
JOURNAL_RC = {
    "figure.dpi": 110, "savefig.dpi": 250,
    "font.family": "DejaVu Sans", "font.size": 9,
    "axes.titlesize": 10, "axes.labelsize": 9,
    "xtick.labelsize": 8, "ytick.labelsize": 8, "legend.fontsize": 8,
    "axes.linewidth": 0.8, "lines.linewidth": 1.0,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.linestyle": ":", "grid.alpha": 0.4,
}


def plot_per_person(pp_50: pd.DataFrame, pp_100: pd.DataFrame, out: str):
    """Bar chart of per-person F-stats, colored by significance."""
    plt.rcParams.update(JOURNAL_RC)
    fig, axes = plt.subplots(1, 2, figsize=(16, 6))

    for ax, df, freq, color in [
        (axes[0], pp_50, "50Hz", "#4477AA"),
        (axes[1], pp_100, "100Hz", "#EE6677"),
    ]:
        names = df["Name"].tolist()
        f_vals = df["F_stat"].tolist()
        sigs = df["Significant"].tolist()
        x = np.arange(len(names))
        colors = [color if s == "Yes" else "#BDBDBD" for s in sigs]
        ax.bar(x, f_vals, color=colors, edgecolor="black", lw=0.5)
        ax.set_xticks(x)
        ax.set_xticklabels(names, rotation=45, ha="right", fontsize=7)
        ax.set_ylabel("F-statistic")
        n_sig = sum(1 for s in sigs if s == "Yes")
        ax.set_title(f"{freq} — Per-person ANOVA F-stats  "
                     f"({n_sig}/{len(names)} significant)", loc="left")
        from matplotlib.patches import Patch
        ax.legend(handles=[
            Patch(facecolor=color, edgecolor="black", label="Significant (p<0.05)"),
            Patch(facecolor="#BDBDBD", edgecolor="black", label="Not significant"),
        ], fontsize=7)

    fig.suptitle("Per-Person 1-Way ANOVA: SBF across 3 time segments", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(out, dpi=250, bbox_inches="tight")
    plt.close(fig)


def plot_segments(data: dict, out: str):
    """Bar chart of combined group means for 3 segments, both freqs."""
    plt.rcParams.update(JOURNAL_RC)
    fig, axes = plt.subplots(1, 2, figsize=(12, 5))

    for ax, freq, color in [(axes[0], 50, "#4477AA"), (axes[1], 100, "#EE6677")]:
        seg_means_all = [[], [], []]
        for name in sorted(data.keys()):
            if freq not in data[name]:
                continue
            for i, m in enumerate(data[name][freq]["means"]):
                seg_means_all[i].append(m)

        if not seg_means_all[0]:
            ax.set_title(f"{freq} Hz — no data"); continue

        group_means = [np.mean(s) for s in seg_means_all]
        group_stds = [np.std(s, ddof=1) for s in seg_means_all]
        group_sems = [std / np.sqrt(len(s)) for std, s in zip(group_stds, seg_means_all)]

        x = np.arange(3)
        ax.bar(x, group_means, yerr=group_sems, capsize=5,
               color=color, edgecolor="black", lw=0.5, alpha=0.85)
        ax.set_xticks(x)
        ax.set_xticklabels(SEG_LABELS)
        ax.set_ylabel("Group mean SBF (perfusion units) +/- SEM")
        ax.set_title(f"{freq} Hz — Group segment means (n={len(seg_means_all[0])})", loc="left")

    fig.suptitle("Combined ANOVA: Group means per time segment", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(out, dpi=250, bbox_inches="tight")
    plt.close(fig)


def plot_simulation(sim_df: pd.DataFrame, out: str):
    """Line plot: % significant combos vs number removed."""
    plt.rcParams.update(JOURNAL_RC)
    fig, ax = plt.subplots(figsize=(10, 5))

    summary = []
    for k in range(7):
        subset = sim_df[sim_df["n_removed"] == k]
        n_total = len(subset)
        if n_total == 0:
            continue
        pct_50 = 100 * (subset["sig_50hz"] == "Yes").sum() / n_total
        pct_100 = 100 * (subset["sig_100hz"] == "Yes").sum() / n_total
        summary.append({"k": k, "pct_sig_50hz": pct_50, "pct_sig_100hz": pct_100,
                        "n_combos": n_total})

    sdf = pd.DataFrame(summary)
    ax.plot(sdf["k"], sdf["pct_sig_50hz"], "o-", color="#4477AA", lw=2, ms=8,
            label="50 Hz")
    ax.plot(sdf["k"], sdf["pct_sig_100hz"], "s-", color="#EE6677", lw=2, ms=8,
            label="100 Hz")

    for _, row in sdf.iterrows():
        ax.annotate(f"{row['pct_sig_50hz']:.1f}%", (row['k'], row['pct_sig_50hz']),
                    textcoords="offset points", xytext=(0, 10), ha="center",
                    fontsize=7, color="#4477AA")
        ax.annotate(f"{row['pct_sig_100hz']:.1f}%", (row['k'], row['pct_sig_100hz']),
                    textcoords="offset points", xytext=(0, -14), ha="center",
                    fontsize=7, color="#EE6677")
        ax.annotate(f"({row['n_combos']:,})", (row['k'], 0),
                    textcoords="offset points", xytext=(0, 5), ha="center",
                    fontsize=6, color="#999")

    ax.set_xlabel("Number of subjects removed")
    ax.set_ylabel("% of combinations with significant ANOVA (p<0.05)")
    ax.set_title("Simulation: How robust is significance to subject removal?",
                 loc="left", fontsize=11)
    ax.set_xticks(range(7))
    ax.set_ylim(-5, 105)
    ax.legend(fontsize=9)
    ax.axhline(50, color="grey", ls="--", lw=0.6, alpha=0.5)

    fig.tight_layout()
    fig.savefig(out, dpi=250, bbox_inches="tight")
    plt.close(fig)


# ══════════════════════════════════════════════════════════════════════════
# 10. CLI
# ══════════════════════════════════════════════════════════════════════════
def main():
    p = argparse.ArgumentParser(
        description="1-Way ANOVA simulation on 26 subjects' SBF after data")
    p.add_argument("--dir", default=".")
    p.add_argument("--excel", default="flux data 20 subjects.xlsx")
    p.add_argument("--sheets", nargs="+",
                   default=["12 subjects", "+ 8 subjects"])
    p.add_argument("--fs", type=float, default=40.0)
    p.add_argument("--hampel-window", type=int, default=21)
    p.add_argument("--hampel-sigmas", type=float, default=3.0)
    p.add_argument("--savgol-window", type=int, default=41)
    p.add_argument("--no-hampel", action="store_true")
    p.add_argument("--no-savgol", action="store_true")
    p.add_argument("--max-remove", type=int, default=6)
    p.add_argument("--out-prefix", default=None)
    args = p.parse_args()

    if args.out_prefix is None:
        args.out_prefix = os.path.join(args.dir, "anova_simulation")

    cfg = SBFConfig(
        sampling_rate_hz=args.fs,
        use_hampel=not args.no_hampel, use_savgol=not args.no_savgol,
        hampel_window=args.hampel_window, hampel_n_sigmas=args.hampel_sigmas,
        savgol_window=args.savgol_window,
    )

    # Load subjects
    print("=== Loading subjects ===")
    excel_path = os.path.join(args.dir, args.excel) if not os.path.isabs(args.excel) else args.excel
    excel_subjects = load_excel_subjects(excel_path, args.sheets)
    txt_subjects = load_txt_subjects(args.dir)
    all_subjects = excel_subjects + txt_subjects
    print(f"  Total: {len(all_subjects)} subjects\n")

    # Prepare data (filter + extract segments + compute means)
    print("=== Preparing data (filtering + segment extraction) ===")
    data = prepare_data(all_subjects, cfg)
    valid = [n for n in data if data[n]]
    print(f"  {len(valid)} subjects with valid segment data\n")

    # Part A: Per-person ANOVA
    print("=== Part A: Per-person ANOVA ===")
    pp_50 = per_person_anova(data, 50)
    pp_100 = per_person_anova(data, 100)
    n_sig_50 = (pp_50["Significant"] == "Yes").sum()
    n_sig_100 = (pp_100["Significant"] == "Yes").sum()
    print(f"  50Hz: {n_sig_50}/{len(pp_50)} significant")
    print(f"  100Hz: {n_sig_100}/{len(pp_100)} significant\n")

    # Part B: Combined ANOVA
    print("=== Part B: Combined ANOVA ===")
    comb_50 = combined_anova(data, 50)
    comb_100 = combined_anova(data, 100)
    print(f"  50Hz:  F={comb_50['F']:.4f}, p={comb_50['p']:.6f}, "
          f"eta²={comb_50['eta_sq']}, {'SIG' if comb_50['sig'] else 'ns'}")
    print(f"  100Hz: F={comb_100['F']:.4f}, p={comb_100['p']:.6f}, "
          f"eta²={comb_100['eta_sq']}, {'SIG' if comb_100['sig'] else 'ns'}")
    if comb_50["tukey"]:
        print(f"  50Hz Tukey: {'; '.join(comb_50['tukey'])}")
    if comb_100["tukey"]:
        print(f"  100Hz Tukey: {'; '.join(comb_100['tukey'])}")
    print()

    # Part C: Simulation
    print("=== Part C: Simulation loop ===")
    sim_df = run_simulation(data, max_remove=args.max_remove)
    print()

    # Influence analysis
    infl_df = influence_analysis(sim_df, comb_50["p"], comb_100["p"])

    # Simulation summary
    print("=== Simulation Summary ===")
    for k in range(args.max_remove + 1):
        subset = sim_df[sim_df["n_removed"] == k]
        n_total = len(subset)
        if n_total == 0:
            continue
        pct_50 = 100 * (subset["sig_50hz"] == "Yes").sum() / n_total
        pct_100 = 100 * (subset["sig_100hz"] == "Yes").sum() / n_total
        print(f"  Remove {k}: {n_total:>7,} combos | "
              f"50Hz sig: {pct_50:5.1f}% | 100Hz sig: {pct_100:5.1f}%")

    # Export
    print("\n=== Writing outputs ===")
    os.makedirs(os.path.dirname(args.out_prefix) or ".", exist_ok=True)
    xlsx_out = args.out_prefix + "_results.xlsx"
    export_excel(pp_50, pp_100, comb_50, comb_100, sim_df, infl_df, cfg, xlsx_out)

    print("  Generating plots ...")
    plot_per_person(pp_50, pp_100, args.out_prefix + "_01_per_person.png")
    plot_segments(data, args.out_prefix + "_02_segments.png")
    plot_simulation(sim_df, args.out_prefix + "_03_simulation.png")

    print(f"\nWrote:")
    print(f"  {xlsx_out}")
    print(f"  {args.out_prefix}_01_per_person.png")
    print(f"  {args.out_prefix}_02_segments.png")
    print(f"  {args.out_prefix}_03_simulation.png")
    print("\nDone!")


if __name__ == "__main__":
    main()
